"""Belge tanıma: dosya adına ve klasör düzenine bakmadan, İÇERİKTEN (PROJE_TALIMATI.md §13).

Her dosyanın ilk ~1500 karakteri okunur (PDF 1. sayfa, Excel ilk 12 satır, Word ilk paragraflar / tablolar) ve
`KURALLAR` ile türü bulunur. Kurallar `src/tani.py`'den (5 firmanın 455 dosyasında denenmiş) aynen alınmıştır.

Bir türden birden fazla dosya varsa hangisinin kullanılacağı yine içerikten seçilir (beyannamenin dönemi, listenin dönem
sütunu, tutarların beyanla tutması); içerikten karar verilemezse kullanıcıya sorulur (`BelirsizSecim`).
Kullanıcı bir dosyanın türünü elle düzeltebilir; düzeltme firma klasöründeki `belge_turleri.yaml`'a yazılır ve
sonraki aylarda (dosya adındaki ay/yıl değişse de) hatırlanır.
"""
import collections
import hashlib
import json
import os
import re
import subprocess
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from .ortak import AYLAR, Donem, excel_satirlari, katla, tr

ILK = 1500
# (tür, açıklama, regex) — sıra önemli: ilk eşleşen kazanır (src/tani.py ile aynı)
KURALLAR = [
    ('KDV1', '1 No.lu KDV beyannamesi', r'KATMA DEĞER VERGİSİ BEYANNAMESİ.*Gerçek Usulde'),
    ('KDV2', '2 No.lu KDV beyannamesi (sorumlu)', r'KATMA DEĞER VERGİSİ BEYANNAMESİ.*Vergi Sorumluları'),
    ('GIB_KONTROL', 'GİB KDV İadesi Kontrol Raporu (Özet)', r'KDV Iadesi Kontrol Raporu'),
    ('LISTE_INDIRILECEK', 'İndirilecek KDV listesi', r'(?i)indirilecek kdv list|Toplam İndirilen KDV|Tevkifata Tabi Olmayan Ve Bu Dönemde İndirilen'),
    ('LISTE_YUKLENILEN', 'Yüklenilen KDV listesi / tutanak çalışması', r'(?i)yüklenilen kdv list|Bünyeye Giren'),
    ('LISTE_INDIRILECEK', 'İndirilecek KDV listesi', r'Alış Faturasının Tarihi.*Satıcının'),
    ('LISTE_TEVKIFATLI', 'Tevkifatlı satış faturaları listesi', r'KISMİ TEVKİFAT|TEVKİFAT UYGULAMASI'),
    ('LISTE_GCB', 'Gümrük çıkış beyannameleri listesi', r'GÜMRÜK ÇIKIŞ BEYANNAMELERİ LİSTESİ'),
    ('LISTE_IHRAC_KAYITLI', 'İhraç kayıtlı satış faturaları listesi',
     r'İHRAÇ KAYITLI SATIŞ FATURASI LİSTESİ|İHRAÇ KAYITLI SATIŞLAR|DİİB KAPSAMINDA SATIŞ FATURA|DAHİLDE İŞLEME İZİN BELGESİ KA'),
    ('LISTE_13MADDE', '13. madde satış listesi', r'MADDE KAPSAMINDA YAPILAN SATIŞLARA'),
    ('LISTE_SATIS', 'Satış faturası listesi', r'SATIŞ FATURASI LİSTESİ|SATIŞ FATURALARI LİSTESİ'),
    ('TAKIP', 'Karşıt inceleme takip listesi (firma/muhasebeci)', r'FİRMA\s.*(KDV\s)?AÇIKLAMA SMMM YMM'),
    ('DIIB_TAKIP', 'DİİB takip listesi', r'DAHİLDE İŞLEME İZİ'),
    ('DIIB_HESAP', 'DİİB hesap tablosu', r'KDV Ödenmeksizin Temin Edilen Mal Bedeli'),
    ('SARFIYAT', 'DİİB sarfiyat tablosu', r'Toplam İhtiyaç İthalat Talebi'),
    ('KUR_FARKI', 'Kur farkı tablosu', r'KUR FARKI|Fatura Tutar Bilgi Cari Ünvan'),
    ('SGK_LISTE', 'SGK/işçi listesi', r'(?i)\bssk\b|\bsgk\b|FİRMA ADI: .*FİRMA ADI:'),
    ('IPTAL_DILEKCE', 'İptal / istisna iptali dilekçesi', r'(?i)iptal'),
    ('IMALAT', 'İmalat tablosu', r'İMALAT TABLOSU'),
    ('MIZAN', 'KDV mizanı / muavin', r'KDV MİZAN|M U A V İ N'),
    ('RAPOR', 'Teminat çözüm / tasdik raporu', r'GENEL BİLGİ.{0,20}Raporun amacı'),
    ('KIT', 'Karşıt inceleme tutanağı', r'KARŞIT İNCELEME TUTANAĞI'),
    ('KIT_DEVAM', 'Tutanak devam sayfası (alt firmalar)', r'ALT FİRMALARIN DÖK'),
    ('YMM_YAZISI', 'YMM bilgi isteme yazısı', r'Konu\s*:?\s*Bilgi İsteme'),
    ('EKSIKLIK', 'Eksiklik yazısı', r'(?i)eksik tamamlama|eksiklik|Konu\s*:\s*Eksik'),
    ('TEMINAT_DILEKCE', 'Teminat mektubu kabul dilekçesi',
     r'(?i)teminat mektub.*(kabul|talep edilmiştir)|Artırımlı Temin|İndirimli Temin|teminat kabul'),
    ('RAPOR_KABUL', 'Rapor/ek kabul dilekçesi', r'Sayı\s*:\s*YMM.*VERGİ DAİRESİ MÜDÜRLÜĞÜ'),
    ('MAHSUP_DILEKCE', 'Mahsup dilekçesi', r'(?i)mahsup'),
    ('SOZLESME', 'Tasdik sözleşmesi', r'SÖZLEŞME|YMM VE BAĞ'),
    ('FATURA_DOKUM', 'Firma fatura dökümü', r'Faturanın Tarihi Faturanın Serisi'),
]
EK_TURLER = {
    'TARANMIS': 'Metni olmayan (taranmış) PDF / görüntü',
    'OKUNAMADI': 'Dosya açılamadı',
    'BILINMEYEN': 'Tanınamadı',
    'DESTEKLENMEYEN': 'Okunmayan dosya türü',
    'YOK_SAY': 'Kullanıcı: yok sayılsın',
}
ACIKLAMA = {**{k: a for k, a, _ in KURALLAR}, **EK_TURLER}
TUR_LISTESI = list(dict.fromkeys([k for k, _, _ in KURALLAR] + list(EK_TURLER)))
KURAL_SURUMU = hashlib.sha1(json.dumps(KURALLAR, ensure_ascii=False).encode()).hexdigest()[:10]

# Raporda kullanılan roller (sıra = tabloda gösterim sırası)
ROLLER = {
    'KDV1': ('1 No.lu KDV beyannamesi', True),
    'LISTE_INDIRILECEK': ('İndirilecek KDV listesi', True),
    'TAKIP': ('Karşıt inceleme takip listesi', True),
    'TEMINAT_DILEKCE': ('Teminat mektubu kabul dilekçesi', True),
    'LISTE_YUKLENILEN': ('Yüklenilen KDV tutanak çalışması', False),
}
BILGI_TURLERI = ['KDV2', 'LISTE_TEVKIFATLI', 'LISTE_IHRAC_KAYITLI', 'LISTE_GCB', 'IMALAT', 'DIIB_TAKIP', 'DIIB_HESAP', 'GIB_KONTROL']

GORUNTU = {'.jpg', '.jpeg', '.png', '.tif', '.tiff', '.bmp', '.gif'}
OKUNAN = {'.pdf', '.xls', '.xlsx', '.xlsm', '.docx', '.doc'}
ATLANAN_DOSYALAR = {'thumbs.db', 'desktop.ini', '.ds_store'}
ATLANAN_KLASORLER = {'CLAUDE TASLAK', 'CLAUDE OUTPUTS', '_CEVRILEN', 'GERIYE DONUK TEST', '__MACOSX'}
COK_DOSYA = 3000           # bir "ay klasörü" için makul üst sınır
CIKTI_ISARETI = '.teminat_cikti'   # programın çıktı klasörlerine konur; bu klasörler hiçbir aramada girdi sayılmaz


class BelirsizSecim(Exception):
    """İçerikten karar verilemedi; kullanıcı seçmeli."""

    def __init__(self, rol, adaylar, aciklama):
        super().__init__(aciklama)
        self.rol, self.adaylar, self.aciklama = rol, list(adaylar), aciklama


class TanimaHatasi(Exception):
    pass


def nf(s):
    return unicodedata.normalize('NFC', s)


def atlanir(ad):
    a = ad.lower()
    return ad.startswith(('~', '._', '.~lock')) or a in ATLANAN_DOSYALAR or a.endswith('.tmp') or a == CIKTI_ISARETI


def atlanan_klasor(ad):
    return katla(ad).upper() in ATLANAN_KLASORLER or ad.startswith('.')


def cikti_klasoru_mu(d):
    """Programın kendi çıktı klasörü mü (adı ne olursa olsun; --cikti ile başka ad verilmiş olabilir)."""
    return (Path(d) / CIKTI_ISARETI).exists()


def cikti_klasoru_hazirla(d):
    """Çıktı klasörünü oluşturur ve işaretler: içindeki taslaklar sonraki taramalarda girdi ya da şablon sanılmaz."""
    d = Path(d)
    d.mkdir(parents=True, exist_ok=True)
    try:
        (d / CIKTI_ISARETI).touch(exist_ok=True)
    except OSError:
        pass
    return d


def icerik_ozeti(p):
    """Dosyanın içerik özeti (aynı boyutlu ama farklı içerikli dosyaları ayırmak için)."""
    h = hashlib.sha1()
    with open(p, 'rb') as f:
        for parca in iter(lambda: f.read(1 << 20), b''):
            h.update(parca)
    return h.hexdigest()


def ayni_dosya(a, b):
    try:
        return os.path.normcase(str(Path(a).resolve())) == os.path.normcase(str(Path(b).resolve()))
    except OSError:
        return False


def dosyalari_listele(klasor, en_cok=COK_DOSYA):
    klasor = Path(klasor)
    out = []
    for kok, dirs, files in os.walk(klasor):
        if cikti_klasoru_mu(kok):
            dirs[:] = []
            continue
        dirs[:] = sorted(d for d in dirs if not atlanan_klasor(d))
        for f in sorted(files):
            if not atlanir(f):
                out.append(Path(kok) / f)
                if len(out) > en_cok:
                    raise TanimaHatasi(f'{klasor} içinde {en_cok}\'den fazla dosya var; bu bir ayın klasörü gibi görünmüyor. '
                                       'Firmanın o ayki klasörünü seçin.')
    return out


# ---------------------------------------------------------------- önbellek (yalnızca tür ve dönem saklanır, metin saklanmaz)
class Onbellek:
    def __init__(self, yol=None):
        self.yol = Path(yol) if yol else Path.home() / '.teminat_cozum' / 'tanima_onbellek.json'
        try:
            self.veri = json.loads(self.yol.read_text(encoding='utf-8'))
        except (OSError, ValueError):
            self.veri = {}
        self.degisti = False

    @staticmethod
    def anahtar(p, ek=''):
        try:
            st = Path(p).stat()
        except OSError:
            return None
        return f'{KURAL_SURUMU}|{ek}|{Path(p).resolve()}|{st.st_size}|{st.st_mtime_ns}'

    def al(self, p, ek=''):
        a = self.anahtar(p, ek)
        return self.veri.get(a) if a else None

    def koy(self, p, deger, ek=''):
        a = self.anahtar(p, ek)
        if a:
            self.veri[a] = deger
            self.degisti = True

    def kaydet(self):
        if not self.degisti:
            return
        try:
            self.yol.parent.mkdir(parents=True, exist_ok=True)
            if len(self.veri) > 50000:
                self.veri = dict(list(self.veri.items())[-30000:])
            self.yol.write_text(json.dumps(self.veri, ensure_ascii=False), encoding='utf-8')
        except OSError:
            pass


class _Hafiza(Onbellek):
    def __init__(self):
        self.yol, self.veri, self.degisti = None, {}, False

    def kaydet(self):
        pass


# ---------------------------------------------------------------- içerik okuma
def _excel_metni(p):
    out = []
    sayfa_sayisi = collections.Counter()
    for sayfa, r in excel_satirlari(p, en_cok=12):
        sayfa_sayisi[sayfa] += 1
        if len(sayfa_sayisi) > 3:
            break
        out.append(' '.join(str(c) for c in r if c is not None and str(c).strip()))
    return '\n'.join(out)[:ILK]


def _docx_metni(p, en_cok=ILK):
    from docx import Document
    d = Document(str(p))
    t = [x.text for x in d.paragraphs]
    for tb in d.tables[:2]:
        for r in tb.rows[:6]:
            t.append(' '.join(c.text for c in r.cells))
    return '\n'.join(t)[:en_cok]


def _pdf_metni(p):
    from .okuyucular import kdv1
    return kdv1.pdf_ilk_sayfa(p)[:ILK]


def doc_metinleri(yollar, zaman_asimi=600, parti_boyu=60):
    """Eski .doc dosyalarını toplu LibreOffice çağrılarıyla metne çevirir: {yol: metin}.

    Her dosya geçici klasöre sıra numarasıyla kopyalanıp öyle çevrilir: farklı klasörlerdeki aynı adlı (ya da yalnızca
    büyük/küçük harfi farklı) dosyaların çıktıları birbirinin üzerine yazılmaz. Orijinallere dokunulmaz.
    """
    import shutil

    from .donusum import soffice_yolu
    so = soffice_yolu()
    out = {}
    yollar = list(yollar)
    if not so or not yollar:
        return out
    for bas in range(0, len(yollar), parti_boyu):
        parti = yollar[bas:bas + parti_boyu]
        with tempfile.TemporaryDirectory() as kaynak, tempfile.TemporaryDirectory() as hedef, \
                tempfile.TemporaryDirectory() as profil:
            kopyalar = []
            for i, p in enumerate(parti):
                k = Path(kaynak) / f'{i}.doc'
                try:
                    shutil.copyfile(p, k)
                except OSError:
                    continue
                kopyalar.append((p, k))
            if not kopyalar:
                continue
            komut = [so, f'-env:UserInstallation={Path(profil).as_uri()}', '--headless', '--convert-to',
                     'txt:Text (encoded):UTF8', '--outdir', hedef, *(str(k) for _, k in kopyalar)]
            try:
                subprocess.run(komut, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=zaman_asimi)
            except (subprocess.TimeoutExpired, OSError):
                pass
            for p, k in kopyalar:
                t = Path(hedef) / (k.stem + '.txt')
                if t.exists():
                    out[p] = t.read_text(encoding='utf-8', errors='ignore')[:ILK]
    return out


def siniflandir(metin):
    t = re.sub(r'\s+', ' ', metin)
    if not t.strip():
        return 'TARANMIS'
    for k, _, rx in KURALLAR:
        if re.search(rx, t):
            return k
    return 'BILINMEYEN'


# ---------------------------------------------------------------- elle tür düzeltmeleri
_YIL_AY = [r'(?<!\d)(?:19|20)\d\d[-_. ]?(?:0[1-9]|1[0-2])(?!\d)',      # 202602, 2026-02
           r'(?<!\d)(?:0[1-9]|1[0-2])[-_. ]?(?:19|20)\d\d(?!\d)',      # 02.2026, 02-2026
           r'(?<!\d)(?:19|20)\d\d(?!\d)']                               # 2026


def ad_anahtari(ad):
    """Bir dosya adının (ya da göreli yolun) aydan bağımsız anahtarı: ay adları ve yıllar çıkarılır, diğer sayılar kalır.

    'liste ŞUBAT 2026.xls' ile 'liste Mart 2027.xls' aynı; 'KDV 1.pdf' ile 'KDV 2.pdf', 'liste.xls' ile 'SİSTEM/liste.xls' farklı.
    """
    parcalar = []
    for parca in str(ad).replace('\\', '/').split('/'):
        n = katla(nf(parca)).upper()
        for d in _YIL_AY:
            n = re.sub(d, ' ', n)
        for a in AYLAR:
            n = re.sub(rf'(?<![A-Z]){katla(a).upper()}(?![A-Z])', ' ', n)
        n = ' '.join(n.split())
        if n:
            parcalar.append(n)
    return '/'.join(parcalar)


class TurDuzeltmeleri:
    """Kullanıcının elle verdiği türler: {ay klasörüne göre göreli yolun anahtarı: TÜR}.

    Firma ayar dosyasının yanındaki belge_turleri.yaml'da saklanır (firmaya özel; ortak bir dosya kullanılmaz).
    """

    def __init__(self, yol=None):
        self.yol = Path(yol) if yol else None
        self.veri = {}
        self.bozuk = None
        if self.yol and self.yol.exists():
            import yaml
            try:
                okunan = yaml.safe_load(self.yol.read_text(encoding='utf-8')) or {}
                if not isinstance(okunan, dict):
                    raise ValueError('sözlük değil')
                self.veri = {str(k): str(v) for k, v in okunan.items()}
            except Exception as e:  # noqa: BLE001 - bozuk dosya: düzeltmesiz devam, ama üzerine yazılmaz
                self.veri = {}
                self.bozuk = f'{self.yol} okunamadı ({type(e).__name__}); elle tür düzeltmeleri uygulanmadı. Dosyayı düzeltin.'

    @staticmethod
    def anahtar(p, kok=None):
        p = Path(p)
        try:
            rel = p.relative_to(kok).as_posix() if kok else p.name
        except ValueError:
            rel = p.name
        return ad_anahtari(rel)

    def tur(self, p, kok=None):
        return self.veri.get(self.anahtar(p, kok))

    def _yaz(self):
        if self.bozuk:
            raise ValueError(self.bozuk)
        if not self.yol:
            raise ValueError('Tür düzeltmesini kaydetmek için önce firmanın ayar dosyasını seçin '
                             '(düzeltmeler o firmaya özel olarak ayar dosyasının yanında saklanır).')
        import yaml
        self.yol.parent.mkdir(parents=True, exist_ok=True)
        bas = ('# Belge türü düzeltmeleri (programın pencereden kaydettiği). Anahtar: ayın klasörüne göre göreli yol; ay adları\n'
               '# ve yıllar çıkarılmıştır, böylece sonraki aylarda da geçerli olur. Tür adları: ' + ', '.join(TUR_LISTESI) + '\n')
        gecici = self.yol.with_name(self.yol.name + '.yeni')
        gecici.write_text(bas + yaml.safe_dump(self.veri, allow_unicode=True, sort_keys=True), encoding='utf-8')
        os.replace(gecici, self.yol)

    def ayarla(self, p, tur, kok=None):
        if tur not in TUR_LISTESI:
            raise ValueError(f'Bilinmeyen tür: {tur}')
        if self.bozuk:
            raise ValueError(self.bozuk)
        self.veri[self.anahtar(p, kok)] = tur
        self._yaz()

    def kaldir(self, p, kok=None):
        """Elle düzeltmeyi kaldırır (tür yeniden içerikten belirlenir)."""
        if self.veri.pop(self.anahtar(p, kok), None) is not None:
            self._yaz()


# ---------------------------------------------------------------- tanıma
@dataclass
class Belge:
    yol: Path
    goreli: str
    tur: str
    elle: bool = False
    donem: Donem = None
    rol: str = ''            # raporda kullanılıyorsa rol adı
    notu: str = ''

    @property
    def aciklama(self):
        return ACIKLAMA.get(self.tur, '')


def tani(klasor, duzeltmeler=None, onbellek=None, ilerleme=None):
    """Klasördeki (alt klasörler dahil) bütün dosyaları içerikten tanır: [Belge]."""
    klasor = Path(klasor)
    onbellek = onbellek if onbellek is not None else Onbellek()
    duzeltmeler = duzeltmeler or TurDuzeltmeleri()
    yollar = dosyalari_listele(klasor)
    belgeler, doc_bekleyen = [], []
    for i, p in enumerate(yollar):
        if ilerleme and i % 50 == 0:
            ilerleme(f'{i}/{len(yollar)} dosya okunuyor…')
        b = Belge(p, nf(p.relative_to(klasor).as_posix()), 'BILINMEYEN')
        belgeler.append(b)
        elle = duzeltmeler.tur(p, klasor)
        if elle:
            b.tur, b.elle = elle, True
            continue
        c = onbellek.al(p)
        if c:
            b.tur = c['tur']
            continue
        uz = p.suffix.lower()
        if uz in GORUNTU:
            b.tur = 'TARANMIS'
        elif uz not in OKUNAN:
            b.tur = 'DESTEKLENMEYEN'
        elif uz == '.doc':
            doc_bekleyen.append(b)
            continue
        else:
            try:
                if uz == '.pdf':
                    metin = _pdf_metni(p)
                elif uz in ('.xls', '.xlsx', '.xlsm'):
                    metin = _excel_metni(p)
                else:
                    metin = _docx_metni(p)
                b.tur = siniflandir(nf(metin))
            except Exception as e:  # noqa: BLE001
                b.tur, b.notu = 'OKUNAMADI', f'{type(e).__name__}: {str(e)[:80]}'
        if b.tur != 'OKUNAMADI':
            onbellek.koy(p, {'tur': b.tur})
    if doc_bekleyen:
        if ilerleme:
            ilerleme(f'{len(doc_bekleyen)} eski Word (.doc) dosyası LibreOffice ile okunuyor…')
        metinler = doc_metinleri([b.yol for b in doc_bekleyen])
        for b in doc_bekleyen:
            if b.yol in metinler:
                b.tur = siniflandir(nf(metinler[b.yol]))
                onbellek.koy(b.yol, {'tur': b.tur})
            else:
                b.tur, b.notu = 'OKUNAMADI', 'LibreOffice ile açılamadı (LibreOffice kurulu mu?)'
    onbellek.kaydet()
    return belgeler


# ---------------------------------------------------------------- dönem okuma yardımcıları
_AY_RX = '|'.join(katla(a).upper() for a in AYLAR)


def metindeki_donemler(metin):
    """Metinde geçen 'ŞUBAT/2026', 'Şubat 2026', '2026/Şubat', '02/2026' dönemleri (Counter).

    Tarihler dönem sayılmaz: '16/02/2026' ve '20 Şubat 2026' (önceki ayın dilekçesi sonraki ay verilmiş olabilir)."""
    n = katla(metin).upper()
    c = collections.Counter()
    for a, y in re.findall(rf'(?<![A-Z])(?<!\d )(?<!\d\.)(?<!\d)({_AY_RX})\s*[-/., ]\s*(20\d\d)(?!\d)', n):
        c[Donem(int(y), [katla(x).upper() for x in AYLAR].index(a) + 1)] += 1
    for y, a in re.findall(rf'(?<!\d)(20\d\d)\s*[-/., ]\s*({_AY_RX})(?![A-Z])', n):
        c[Donem(int(y), [katla(x).upper() for x in AYLAR].index(a) + 1)] += 1
    for a, y in re.findall(r'(?<![\d./-])(0?[1-9]|1[0-2])\s*/\s*(20\d\d)(?!\d)', n):
        c[Donem(int(y), int(a))] += 1
    for y, a in re.findall(r'(?<![\d./-])(20\d\d)\s*/\s*(0?[1-9]|1[0-2])(?![\d./])', n):
        c[Donem(int(y), int(a))] += 1
    return c


def rapor_bilgisi(p, onbellek=None):
    """Teminat çözüm raporunun kapağından dönem ve mükellef hesap no (VKN): {'donem': Donem|None, 'vkn': str|None}."""
    onbellek = onbellek if onbellek is not None else _Hafiza()
    c = onbellek.al(p, 'rapor')
    if c is not None:
        return {'donem': Donem.coz(c['donem']) if c.get('donem') else None, 'vkn': c.get('vkn')}
    from docx import Document

    from .donusum import DonusumHatasi, docx_hazirla
    donem = vkn = None
    try:
        with tempfile.TemporaryDirectory() as tmp:
            d = Document(str(docx_hazirla(p, tmp)))
            for t in d.tables[:2]:
                for r in t.rows:
                    hucre = [c.text.strip() for c in r.cells]
                    satir = ' | '.join(hucre)
                    if donem is None and 'İncelemenin Dönemi' in satir:
                        m = re.search(r'([A-ZÇĞİÖŞÜ]+)-(\d{4})', hucre[-1])
                        if m:
                            try:
                                donem = Donem.coz(f'{m.group(1)}-{m.group(2)}')
                            except ValueError:
                                pass
                    if vkn is None and re.search(r'Hesap No|Vergi Kimlik|VKN', satir) and 'Mükellef' in satir:
                        rakam = re.sub(r'\D', '', hucre[-1])
                        if len(rakam) in (10, 11):
                            vkn = rakam
    except (DonusumHatasi, Exception):  # noqa: BLE001 - okunamayan rapor: bilgi yok
        return {'donem': None, 'vkn': None}
    onbellek.koy(p, {'donem': donem.tire if donem else None, 'vkn': vkn}, 'rapor')
    return {'donem': donem, 'vkn': vkn}


def metindeki_vknler(metin):
    """Metinde geçen 10-11 haneli numaralar (aralarında tek boşluk olabilir: '123 456 7890')."""
    return {re.sub(r'\D', '', x) for x in re.findall(r'(?<![\d])(\d(?:[ ]?\d){9,10})(?![\d])', metin or '')}


def vkn_metinde(vkn, metin):
    return bool(vkn) and vkn in metindeki_vknler(metin)


# ---------------------------------------------------------------- rol seçimi
@dataclass
class Secim:
    klasor: Path
    belgeler: list
    yollar: dict = field(default_factory=dict)        # rol → Path
    donem: Donem = None
    k: dict = None                                      # okunmuş KDV 1
    bu_ayin_raporu: Path = None
    notlar: list = field(default_factory=list)         # [(durum, konu, açıklama)]

    def eksik_zorunlu(self):
        return [r for r, (_, z) in ROLLER.items() if z and r not in self.yollar]


def _adaylar(belgeler, tur):
    return [b for b in belgeler if b.tur == tur]


def _goreli(b):
    return b.goreli


def _imza(veri):
    """Okunan satırların özeti: aynı puanlı iki aday gerçekten aynı içerikte mi?"""
    return hashlib.sha1(json.dumps(veri, sort_keys=True, default=str, ensure_ascii=False).encode('utf-8')).hexdigest()


def _zorla_belgesi(s, tur, yol):
    """Kullanıcının elle seçtiği dosya: tanıma listesindeki kaydı (türü farklıysa düzeltilerek) ya da yeni kayıt."""
    yol = Path(yol)
    for b in s.belgeler:
        if ayni_dosya(b.yol, yol):
            if b.tur != tur:
                b.tur, b.elle = tur, True
            return b
    b = Belge(yol, str(yol), tur, elle=True)
    s.belgeler.append(b)
    return b


def _en_iyiler(puan, anahtar):
    """En yüksek puanlı adaylar ve içerik imzaları farklı mı: (en iyiler, farklı_mı)."""
    ust = [x for x in puan if anahtar(x) == anahtar(puan[0])]
    return ust, len({x[-1] for x in ust}) > 1


def kdv1_sec(adaylar, istenen=None, klasor_ayi=None):
    """Okunabilen 1 No.lu beyannameler arasından seçer: (Belge, k, notlar)."""
    from .okuyucular.kdv1 import kdv1_donem, kdv1_metinden, pdf_metni
    notlar = []
    okunan = []
    for b in adaylar:
        try:
            k = kdv1_metinden(pdf_metni(b.yol), b.yol)
        except Exception as e:  # noqa: BLE001
            b.notu = f'okunamadı: {type(e).__name__}'
            continue
        yd = kdv1_donem(k)
        if yd and ('yurtici_alim' in k or 'matrah_toplam' in k):
            b.donem = Donem(*yd)
            okunan.append((b, k))
        else:
            b.notu = 'beyanname olarak tanındı ama dönem/tutarları okunamadı'
    if not okunan:
        if adaylar:
            raise TanimaHatasi('1 No.lu KDV beyannamesi olarak tanınan dosyalar okunamadı: '
                               + ', '.join(_goreli(b) for b in adaylar) + '. Beyanname PDF\'i metin içermeli (taranmış olmamalı).')
        raise TanimaHatasi('Klasörde 1 No.lu KDV beyannamesi bulunamadı (hiçbir PDF\'te "KATMA DEĞER VERGİSİ BEYANNAMESİ … '
                           'Gerçek Usulde" yazısı yok ya da beyanname elle başka türe çevrilmiş). Beyannameyi klasöre koyun '
                           'ya da tanıma tablosunda türünü düzeltin.')
    if istenen:
        uyan = [x for x in okunan if x[0].donem == istenen]
        if not uyan:
            raise TanimaHatasi(f'Klasörde {istenen.tire} dönemine ait beyanname yok; bulunanlar: '
                               + ', '.join(f'{b.donem.tire} ({_goreli(b)})' for b, _ in okunan))
        okunan = uyan
    donemler = sorted({b.donem for b, _ in okunan})
    if len(donemler) > 1:
        if klasor_ayi and [d for d in donemler if d.ay == klasor_ayi[0] and (klasor_ayi[1] in (None, d.yil))]:
            secilen_d = max(d for d in donemler if d.ay == klasor_ayi[0] and (klasor_ayi[1] in (None, d.yil)))
            neden = 'klasör adındaki ayla uyan'
        else:
            secilen_d = donemler[-1]
            neden = 'en son dönem'
        notlar.append(('UYARI', 'Belge seçimi', f'Klasörde birden fazla dönemin beyannamesi var ({", ".join(d.tire for d in donemler)}); '
                       f'{neden} {secilen_d.tire} seçildi. Başka bir ay için o ayın klasörünü seçin.'))
        okunan = [x for x in okunan if x[0].donem == secilen_d]
    if len(okunan) > 1:
        imza = {(k.get('matrah_toplam'), k.get('iade_gereken'), k.get('sonraki_devreden'), k.get('duzeltme')) for _, k in okunan}
        if len(imza) > 1:
            raise BelirsizSecim('KDV1', [b.yol for b, _ in okunan],
                                f'Aynı dönemin ({okunan[0][0].donem.tire}) içeriği farklı {len(okunan)} beyannamesi var '
                                '(ör. asıl ve düzeltme beyannamesi). Hangisinin kullanılacağını seçin.')
        notlar.append(('BİLGİ', 'Belge seçimi', f'Aynı beyannamenin {len(okunan)} kopyası var; {_goreli(okunan[0][0])} kullanıldı.'))
    return okunan[0][0], okunan[0][1], notlar


def _liste_sec(adaylar, donem, base, notlar, elle=False):
    from .okuyucular.listeler import donem_degeri, indirilecek_oku
    puan = []
    for b in adaylar:
        try:
            rows = indirilecek_oku(b.yol)
        except Exception as e:  # noqa: BLE001
            b.notu = f'okunamadı: {type(e).__name__}'
            continue
        if not rows:
            b.notu = 'satır okunamadı (sütun düzeni farklı olabilir)'
            continue
        donemler = collections.Counter(donem_degeri(r['donem']) for r in rows)
        uygun = donemler.get((donem.yil, donem.ay), 0) / len(rows)
        en_cok = donemler.most_common(1)[0][0]
        if en_cok:
            b.donem = Donem(*en_cok)
        top = sum(r['kdv'] for r in rows)
        b.notu = f'{len(rows)} satır, toplam {tr(top)}'
        puan.append(((round(uygun, 3), -round(abs(top - base), 2)), b, top, _imza(rows)))
    if not puan:
        return None
    puan.sort(key=lambda x: x[0], reverse=True)
    if len(puan) > 1 and not elle:
        ust, farkli = _en_iyiler(puan, lambda x: x[0])
        if farkli:
            raise BelirsizSecim('LISTE_INDIRILECEK', [x[1].yol for x in ust],
                                'Birden fazla indirilecek KDV listesi var; dönem sütunu ve toplamları aynı ama içerikleri farklı. '
                                'Hangisi bu aya ait?')
        neden = 'aynı içerikli kopyalardan ilki' if len(ust) > 1 else f'dönem sütunu {donem.tire} olan ve toplamı beyana en yakın olan'
        notlar.append(('BİLGİ', 'Belge seçimi', f'{len(puan)} indirilecek KDV listesinden {neden} seçildi: {_goreli(puan[0][1])}. Diğerleri: '
                       + ', '.join(f'{_goreli(x[1])} ({x[1].donem.tire if x[1].donem else "dönem ?"}, {tr(x[2])})' for x in puan[1:])))
    return puan[0][1]


def _takip_sec(adaylar, liste_satirlari, notlar, elle=False):
    from .okuyucular.listeler import takip_oku
    from .rapor.safha import firma_anahtari
    by = collections.defaultdict(float)
    for r in liste_satirlari:
        by[firma_anahtari(r['satici'])] += r['kdv']
    tutarlar = {round(v, 2) for v in by.values()}
    puan = []
    for b in adaylar:
        try:
            t = takip_oku(b.yol)
        except Exception as e:  # noqa: BLE001
            b.notu = f'okunamadı: {type(e).__name__}'
            continue
        if not t:
            b.notu = 'firma satırı okunamadı'
            continue
        uyan = sum(1 for x in t if round(x['kdv'], 2) in tutarlar or abs(by.get(firma_anahtari(x['firma']), -1) - x['kdv']) < 0.01)
        b.notu = f'{len(t)} firma, {uyan} tanesi indirilecek listeyle tutuyor'
        puan.append((round(uyan / len(t), 4), b, sum(x['kdv'] for x in t), _imza(t)))
    if not puan:
        return None
    puan.sort(key=lambda x: x[0], reverse=True)
    if len(puan) > 1 and not elle:
        ust, farkli = _en_iyiler(puan, lambda x: x[0])
        if farkli:
            raise BelirsizSecim('TAKIP', [x[1].yol for x in ust],
                                'Birden fazla karşıt inceleme takip listesi var; indirilecek listeyle eşit derecede tutuyorlar ama '
                                'içerikleri farklı (ör. YMM / açıklama sütunu). Hangisi bu aya ait?')
        neden = 'aynı içerikli kopyalardan ilki' if len(ust) > 1 else 'indirilecek listeyle en çok tutan'
        notlar.append(('BİLGİ', 'Belge seçimi', f'{len(puan)} takip listesinden {neden} seçildi: {_goreli(puan[0][1])}.'))
    return puan[0][1]


def dilekce_metni(p):
    from .donusum import docx_hazirla
    with tempfile.TemporaryDirectory() as tmp:
        return _docx_metni(docx_hazirla(p, tmp), en_cok=None)


def _dilekce_sec(adaylar, donem, notlar, elle=False):
    from .okuyucular.teminat import dilekce_metinden
    puan = []
    for b in adaylar:
        try:
            metin = dilekce_metni(b.yol)
        except Exception as e:  # noqa: BLE001
            b.notu = f'okunamadı: {type(e).__name__}'
            continue
        tem = dilekce_metinden(metin)
        donemler = metindeki_donemler(metin)
        if donemler:
            b.donem = donemler.most_common(1)[0][0]
        eslesti = donem in donemler
        p = (2 if eslesti else 0) + (1 if tem.get('tarih') else 0) + (1 if tem.get('toplam') else 0)
        puan.append((p, b, eslesti, _imza(metin)))
    if not puan:
        return None
    puan.sort(key=lambda x: x[0], reverse=True)
    secilen, eslesti = puan[0][1], puan[0][2]
    if len(puan) > 1 and not elle:
        ust, farkli = _en_iyiler(puan, lambda x: x[0])
        if farkli or not eslesti:
            raise BelirsizSecim('TEMINAT_DILEKCE', [x[1].yol for x in (ust if farkli else puan)],
                                f'Birden fazla teminat dilekçesi var ve içeriğinden hangisinin {donem.tire} dönemine ait olduğu '
                                'anlaşılamadı. Hangisi kullanılsın?')
        neden = 'aynı içerikli kopyalardan ilki' if len(ust) > 1 else f'{donem.tire} dönemini anan'
        notlar.append(('BİLGİ', 'Belge seçimi', f'{len(puan)} teminat dilekçesinden {neden} seçildi: {_goreli(secilen)}.'))
    if not eslesti:
        anilan = f'{secilen.donem.tire} dönemini anıyor' if secilen.donem else 'hiçbir dönemi anmıyor'
        notlar.append(('UYARI', 'Belge dönemi', f'Teminat dilekçesi {anilan}, beyanname {donem.tire}: {_goreli(secilen)} — '
                       'önceki ayın dilekçesi olabilir; mektup tarihi/numarası ve tutarları kontrol edin.'))
    return secilen


def _yuklenilen_sec(adaylar, hedef, notlar, elle=False):
    from .okuyucular.listeler import yuklenilen_oku
    puan = []
    for b in adaylar:
        try:
            y = yuklenilen_oku(b.yol)
        except Exception as e:  # noqa: BLE001
            b.notu = f'okunamadı: {type(e).__name__}'
            continue
        if not y:
            b.notu = 'satır okunamadı (tutanak çalışması biçiminde değil)'
            continue
        top = sum(x['kdv'] for x in y)
        b.notu = f'{len(y)} satır, toplam {tr(top)}'
        puan.append((round(abs(top - hedef), 2), b, top, _imza(y)))
    if not puan:
        return None
    puan.sort(key=lambda x: x[0])
    if elle:
        return puan[0][1]
    tutan = [x for x in puan if x[0] <= 0.01]
    if tutan:
        if len({x[-1] for x in tutan}) > 1:
            raise BelirsizSecim('LISTE_YUKLENILEN', [x[1].yol for x in tutan],
                                'Toplamı beyandaki 301 (+339) yüklenilen KDV\'ye eşit birden fazla, içeriği farklı yüklenilen '
                                'listesi var. Hangisi kullanılsın?')
        if len(puan) > 1:
            notlar.append(('BİLGİ', 'Belge seçimi', f'{len(puan)} yüklenilen listesinden toplamı beyana eşit olan seçildi: '
                           f'{_goreli(tutan[0][1])}.'))
        return tutan[0][1]
    if len(puan) > 1:
        raise BelirsizSecim('LISTE_YUKLENILEN', [x[1].yol for x in puan],
                            f'Hiçbir yüklenilen listesinin toplamı beyandaki 301 (+339) yüklenilen KDV\'ye ({tr(hedef)}) eşit değil. '
                            'Hangisi kullanılsın?')
    notlar.append(('UYARI', 'Belge seçimi', f'Yüklenilen tutanak dosyasının toplamı ({tr(puan[0][2])}) beyandaki 301 (+339) '
                   f'yüklenilen KDV\'ye ({tr(hedef)}) eşit değil: {_goreli(puan[0][1])}.'))
    return puan[0][1]


def bu_ayin_raporu_sec(s, onbellek=None):
    """Klasörde bu ayın bitmiş raporu da varsa (taslakla karşılaştırmak için) seçer; aynı mükellefe ait olmayanlar ve
    içerikleri farklı birden fazla aday varsa karşılaştırma yapılmaz (not yazılır)."""
    metin = (s.k or {}).get('_text') or ''
    adaylar = []
    for r in _adaylar(s.belgeler, 'RAPOR'):
        bilgi = rapor_bilgisi(r.yol, onbellek)
        r.donem = bilgi['donem']
        if r.donem != s.donem:
            continue
        if bilgi['vkn'] and metindeki_vknler(metin) and not vkn_metinde(bilgi['vkn'], metin):
            r.notu = 'başka mükellefin raporu (VKN farklı)'
            continue
        adaylar.append(r)
    if not adaylar:
        return
    ozetler = set()
    for r in adaylar:
        try:
            ozetler.add(icerik_ozeti(r.yol))
        except OSError:
            ozetler.add(str(r.yol))
    if len(ozetler) > 1:
        s.notlar.append(('UYARI', 'Gerçek raporla karşılaştırma', f'Klasörde {s.donem.tire} dönemine ait içeriği farklı '
                         f'{len(adaylar)} bitmiş rapor var ({", ".join(_goreli(r) for r in adaylar)}); hangisinin son hali olduğu '
                         'bilinmediği için taslak bunlarla karşılaştırılmadı. Karşılaştırmayı "Gerçek raporla karşılaştır" ile yapın.'))
        return
    s.bu_ayin_raporu = adaylar[0].yol
    adaylar[0].rol = 'BU_AYIN_RAPORU'


def secim_yap(klasor, belgeler, istenen=None, zorla=None, onbellek=None):
    """Tanınan belgelerden raporda kullanılacakları seçer (zorla: {rol: yol} kullanıcı seçimleri)."""
    from .klasorler import ay_coz
    from .okuyucular.listeler import indirilecek_oku
    zorla = zorla or {}
    s = Secim(Path(klasor), belgeler)

    def adaylar(rol):
        if zorla.get(rol):
            return [_zorla_belgesi(s, rol, zorla[rol])], True
        return _adaylar(s.belgeler, rol), False

    klasor_ayi = ay_coz(Path(klasor).name)
    b, k, n = kdv1_sec(adaylar('KDV1')[0], istenen, klasor_ayi)
    s.notlar += n
    b.rol, s.yollar['KDV1'], s.donem, s.k = 'KDV1', b.yol, b.donem, k
    if klasor_ayi and (klasor_ayi[0] != s.donem.ay or (klasor_ayi[1] and klasor_ayi[1] != s.donem.yil)):
        s.notlar.append(('UYARI', 'Belge dönemi', f'Klasör adı ("{Path(klasor).name}") başka bir ayı gösteriyor ama beyanname '
                         f'{s.donem.tire} dönemine ait — doğru klasörü seçtiğinizden emin olun.'))
    base = k.get('yurtici_alim', 0) + k.get('sorumlu', 0) + k.get('ithal', 0)
    a, elle = adaylar('LISTE_INDIRILECEK')
    lb = _liste_sec(a, s.donem, base, s.notlar, elle)
    satirlar = []
    if lb:
        lb.rol, s.yollar['LISTE_INDIRILECEK'] = 'LISTE_INDIRILECEK', lb.yol
        satirlar = indirilecek_oku(lb.yol)
    a, elle = adaylar('TAKIP')
    tb = _takip_sec(a, satirlar, s.notlar, elle)
    if tb:
        tb.rol, s.yollar['TAKIP'] = 'TAKIP', tb.yol
    a, elle = adaylar('TEMINAT_DILEKCE')
    db = _dilekce_sec(a, s.donem, s.notlar, elle)
    if db:
        db.rol, s.yollar['TEMINAT_DILEKCE'] = 'TEMINAT_DILEKCE', db.yol
    hedef = k.get('301_yuklenilen', 0) + k.get('339_iade', 0)
    a, elle = adaylar('LISTE_YUKLENILEN')
    if hedef or elle:
        yb = _yuklenilen_sec(a, hedef, s.notlar, elle)
        if yb:
            yb.rol, s.yollar['LISTE_YUKLENILEN'] = 'LISTE_YUKLENILEN', yb.yol
    elif a:
        s.notlar.append(('BİLGİ', 'Belge seçimi', 'Beyanda 301 yüklenilen KDV yok; yüklenilen listeleri kullanılmadı.'))
    bu_ayin_raporu_sec(s, onbellek)
    for rol, (ad, zorunlu) in ROLLER.items():
        if rol not in s.yollar and zorunlu:
            s.notlar.append(('HATA', 'Eksik belge', f'{ad} bulunamadı (klasörde içeriği bu türe uyan dosya yok). '
                             'Tanıma tablosunda türü yanlış görünen bir dosya varsa türünü düzeltin.'))
    return s


# ---------------------------------------------------------------- şablon (önceki ayın raporu) arama
def sablon_bul(klasor, onceki, kdv1_metni=None, onbellek=None, en_cok=150, ilerleme=None, en_cok_klasor=4000,
               haric=(), elle_rapor=()):
    """Seçilen ay klasörünün çevresindeki Word dosyalarından kapağındaki dönem `onceki` olan teminat çözüm raporunu
    içerikten bulur. Arama yakından uzağa yapılır: önce klasörün kendisi, sonra üstü, sonra onun üstü (her biri 4 seviye
    aşağıya kadar); bir seviyede rapor bulununca daha uzağa bakılmaz. Kapaktaki mükellef VKN'si beyannamede geçmiyorsa
    (başka firmanın raporu) alınmaz. haric: kullanıcının başka türe çevirdiği dosyalar; elle_rapor: kullanıcının RAPOR
    olarak işaretlediği dosyalar. Dönüş: (yol | None, notlar). İçeriği farklı birden çok rapor bulunursa BelirsizSecim."""
    from .klasorler import ay_coz
    klasor = Path(klasor).resolve()
    onbellek = onbellek if onbellek is not None else Onbellek()
    notlar = []
    haric = [Path(x) for x in haric]
    elle_rapor = [Path(x) for x in elle_rapor]
    kokler = [klasor] + list(klasor.parents)[:2]
    gorulen, gezilen_kokler = set(), []
    gezilen = incelenen = 0
    baska_firma, bulunan = [], []

    def oncelik(p):
        yol = katla(str(p)).upper()
        parca = [ay_coz(x) for x in p.parent.parts[-3:]]
        ay_uyar = any(x and x[0] == onceki.ay and x[1] in (None, onceki.yil) for x in parca)
        try:
            boy = p.stat().st_size
        except OSError:
            boy = 0
        return (not ay_uyar, 'RAPOR' not in yol, p.suffix.lower() != '.docx', -boy)

    for sira, kok in enumerate(kokler):
        adaylar = [p for p in elle_rapor if p not in gorulen] if sira == 0 else []
        gorulen.update(adaylar)
        for d, dirs, files in os.walk(kok):
            gezilen += 1
            if gezilen > en_cok_klasor:
                notlar.append(('BİLGİ', 'Şablon arama', f'Çevrede çok fazla klasör var; ilk {en_cok_klasor} klasöre bakıldı.'))
                dirs[:] = []
                break
            dp = Path(d)
            if cikti_klasoru_mu(dp):
                dirs[:] = []
                continue
            rel = len(dp.relative_to(kok).parts)
            dirs[:] = [x for x in dirs if not atlanan_klasor(x) and rel < 4 and (dp / x) not in gezilen_kokler]
            for f in files:
                p = dp / f
                if atlanir(f) or p.suffix.lower() not in ('.doc', '.docx') or p in gorulen:
                    continue
                gorulen.add(p)
                if any(ayni_dosya(p, h) for h in haric):
                    continue
                adaylar.append(p)
        gezilen_kokler.append(kok)
        adaylar.sort(key=lambda p: (p not in elle_rapor, oncelik(p)))
        hak = en_cok - incelenen
        taranacak = adaylar[:hak]
        if len(adaylar) > hak:
            notlar.append(('BİLGİ', 'Şablon arama', f'{kok} çevresindeki {len(adaylar)} Word dosyasından öncelikli {hak} '
                           'tanesine bakıldı.'))
        incelenen += len(taranacak)
        docs = [p for p in taranacak if p.suffix.lower() == '.doc' and onbellek.al(p) is None and p not in elle_rapor]
        if docs and ilerleme:
            ilerleme(f'Önceki ayın raporu aranıyor: {len(docs)} eski Word dosyası okunuyor…')
        metinler = doc_metinleri(docs) if docs else {}
        bulunan = []
        for p in taranacak:
            c = {'tur': 'RAPOR'} if p in elle_rapor else onbellek.al(p)
            if c is None:
                try:
                    metin = metinler.get(p) if p.suffix.lower() == '.doc' else _docx_metni(p)
                except Exception:  # noqa: BLE001
                    metin = None
                if metin is None:
                    continue
                c = {'tur': siniflandir(nf(metin))}
                onbellek.koy(p, c)
            if c['tur'] != 'RAPOR':
                continue
            bilgi = rapor_bilgisi(p, onbellek)
            if bilgi['donem'] == onceki:
                bulunan.append((p, bilgi))
        onbellek.kaydet()
        if kdv1_metni and metindeki_vknler(kdv1_metni):
            ayni = [x for x in bulunan if vkn_metinde(x[1]['vkn'], kdv1_metni)]
            baska_firma += [x for x in bulunan if x[1]['vkn'] and not vkn_metinde(x[1]['vkn'], kdv1_metni)]
            bulunan = ayni or [x for x in bulunan if not x[1]['vkn']]
        if bulunan:
            break
        if incelenen >= en_cok:
            break
    if baska_firma:
        notlar.append(('BİLGİ', 'Şablon arama', f'{len(baska_firma)} rapor başka mükellefe (VKN) ait olduğu için alınmadı: '
                       + ', '.join(str(p) for p, _ in baska_firma[:5])))
    if not bulunan:
        return None, notlar
    if len(bulunan) > 1:
        ozetler = set()
        for p, _ in bulunan:
            try:
                ozetler.add(icerik_ozeti(p))
            except OSError:
                ozetler.add(str(p))
        if len(ozetler) > 1:
            raise BelirsizSecim('SABLON', [p for p, _ in bulunan],
                                f'{onceki.tire} dönemine ait içeriği farklı {len(bulunan)} rapor bulundu. Şablon olarak hangisi '
                                'kullanılsın (ofisin son hali)?')
        notlar.append(('BİLGİ', 'Şablon arama', f'{onceki.tire} raporunun {len(bulunan)} aynı kopyası var; {bulunan[0][0]} kullanıldı.'))
    return bulunan[0][0], notlar


# ---------------------------------------------------------------- tanıma tablosu (metin)
def tablo_metni(secim, sablon=None, sablon_notu=''):
    s = secim
    satirlar = []
    sira = {r: i for i, r in enumerate(list(ROLLER) + ['BU_AYIN_RAPORU'])}
    belgeler = sorted(s.belgeler, key=lambda b: (b.rol == '', sira.get(b.rol, 99), b.tur in EK_TURLER, b.tur, b.goreli))
    gen = max([len(b.goreli) for b in belgeler] + [20])
    gen = min(gen, 70)
    satirlar.append(f'{"TÜR":<20} {"DÖNEM":<12} {"KULLANIM":<16} DOSYA')
    gruplu = collections.Counter()
    for b in belgeler:
        if b.tur in ('BILINMEYEN', 'TARANMIS', 'DESTEKLENMEYEN') and not b.rol and not b.elle:
            gruplu[(b.tur, str(Path(b.goreli).parent))] += 1
            continue
        kullanim = '✔ KULLANILACAK' if b.rol else ''
        if b.rol == 'BU_AYIN_RAPORU':
            kullanim = '✔ karşılaştırma'
        tur = b.tur + (' (elle)' if b.elle else '')
        satirlar.append(f'{tur:<20} {(b.donem.tire if b.donem else ""):<12} {kullanim:<16} {b.goreli}'
                        + (f'   [{b.notu}]' if b.notu else ''))
    for (tur, kl), n in sorted(gruplu.items()):
        satirlar.append(f'{tur:<20} {"":<12} {"":<16} {n} dosya — {kl if kl != "." else "(ana klasör)"}')
    satirlar.append('')
    satirlar.append('--- RAPOR İÇİN DURUM ---')
    satirlar.append(f'Dönem (beyannameden): {s.donem.tire if s.donem else "?"}')
    var = collections.Counter(b.tur for b in s.belgeler)
    for rol, (ad, zorunlu) in ROLLER.items():
        if rol in s.yollar:
            ek = f'  ({var[rol]} dosya arasından seçildi)' if var[rol] > 1 else ''
            satirlar.append(f'✔ {ad}: {Path(s.yollar[rol]).name}{ek}')
        elif zorunlu:
            satirlar.append(f'✘ EKSİK (zorunlu) {ad}')
        elif var[rol]:
            satirlar.append(f'– {ad}: kullanılmadı')
    if sablon:
        satirlar.append(f'✔ Şablon (önceki ayın raporu, içerikten): {sablon}')
    else:
        satirlar.append(f'✘ Şablon (önceki ayın raporu) bulunamadı{": " + sablon_notu if sablon_notu else ""}')
    if s.bu_ayin_raporu:
        satirlar.append(f'✔ Bu ayın bitmiş raporu da var; taslak onunla karşılaştırılacak: {Path(s.bu_ayin_raporu).name}')
    for t in BILGI_TURLERI:
        if var[t]:
            satirlar.append(f'· {ACIKLAMA[t]}: {var[t]} dosya (şimdilik okunmuyor)')
    if var['TARANMIS']:
        satirlar.append(f'! {var["TARANMIS"]} taranmış PDF / görüntü var (metin yok) — gerekiyorsa elle okunmalı')
    if var['BILINMEYEN']:
        satirlar.append(f'? {var["BILINMEYEN"]} dosya tanınamadı (çoğu fatura / GÇB olabilir)')
    if var['OKUNAMADI']:
        satirlar.append(f'! {var["OKUNAMADI"]} dosya açılamadı')
    return '\n'.join(satirlar)
