"""Önceki ayın bitmiş Word raporundan yeni ayın taslağını üretir.

Şablon = firmanın bir önceki ayın raporu. Program yeni rapor yazmaz; önceki raporu kopyalar, dönem metinlerini
ve rakamları günceller, tablolarda satır ekler/siler. Bilinmeyen bilgi [DOLDURULACAK] + sarı; takdir isteyen
her seçim ve her tutarsızlık kontrol listesine yazılır. Şablon dosyasına yazılmaz.
"""
import collections
import copy
import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.table import _Row

from ..kontroller import belge, kurallar
from ..kontroller.liste import KontrolListesi
from ..okuyucular.kdv1 import IADE_ETIKETLERI, iade_kalemleri
from ..ortak import AYLAR, DOLD, PARA_RE, Donem, katla, num, pct, pct2, tr
from . import docx_araclari as dx
from .docx_araclari import (SARI, SablonYok, ara, del_row, find_row, find_rows, grubu_esitle, hucre_isaretle,
                            ilk_hucre, isaretle, metin_degistir, paragraf_degistir, row_text, set_cell, set_row_vals,
                            uniq_cells)
from .safha import firma_anahtari, safha_hesapla
from .tablolar import tablolari_bul

ASGARI_ORAN = 0.80   # karşıt incelemede asgari oran (SMMM-YMM tebliğleri)
IADE_KODLARI = ['301', '318', '339', '350', '410', '448']


class SablonHatasi(Exception):
    """Şablon bu haliyle kullanılamıyor (ör. dönemi okunamadı)."""


@dataclass
class Sonuc:
    cikti: Path
    kontroller: KontrolListesi
    ozet: dict = field(default_factory=dict)


def _not_aktar(kl, notlar):
    for durum, konu, ac in notlar:
        kl.ekle(durum, konu, ac)


class TaslakOlusturucu:
    def __init__(self, sablon_yolu, girdiler, ayar, kati=False):
        self.sablon_yolu = Path(sablon_yolu)
        self.d = Document(str(sablon_yolu))
        self.g = girdiler
        self.k = girdiler.k
        self.ayar = ayar
        self.kati = kati
        self.kl = KontrolListesi()
        self.ozet = {}
        self.yeni = girdiler.donem
        self.T = {}

    # ------------------------------------------------------------ çatı
    def adim(self, ad, fn):
        try:
            fn()
        except SablonYok as e:
            if self.kati:
                raise
            self.kl.hata('Şablon yapısı', f'{ad}: {e} — bu bölüm güncellenemedi, elle doldurun.')
        except Exception as e:  # noqa: BLE001 - taslak yine de üretilsin, hata listeye yazılsın
            if self.kati:
                raise
            self.kl.hata('İşlem hatası', f'{ad}: {type(e).__name__}: {e} — bu bölüm güncellenemedi, elle kontrol edin.')

    def olustur(self, cikti_yolu):
        cikti_yolu = Path(cikti_yolu)
        if cikti_yolu.resolve() == self.sablon_yolu.resolve():
            raise SablonHatasi('Çıktı dosyası şablonun üzerine yazılamaz.')
        dx.kayit_baslat()
        _not_aktar(self.kl, self.g.notlar)
        self.T, sorunlar = tablolari_bul(self.d, self.ayar)
        for ad, ac in sorunlar:
            self.kl.hata('Tablo bulma', ac)
        self._sablon_donemi()
        self.k_kalem = iade_kalemleri(self.k)
        self.base = self.k['yurtici_alim'] + self.k.get('sorumlu', 0) + self.k.get('ithal', 0)
        self.ozet.update(donem=self.yeni.tire, sablon_donemi=self.eski.tire, sablon=self.sablon_yolu.name,
                         indirilecek_beyan=self.base)
        for ad, fn in [('Ocak raporu atıfları', self.ocak_paragraflari),
                       ('Dönem metinleri', self.donem_metinleri),
                       ('Teminat mektubu paragrafları', self.teminat_paragraflari),
                       ('Kapak', self.kapak),
                       ('1-6 İşçi sayısı', self.isci),
                       ('3-1 Matrah tablosu', lambda: self.matrah_doldur('matrah')),
                       ('3-2 İndirimler', self.indirimler),
                       ('3-3-1 İş hacmi', self.is_hacmi),
                       ('3-4 Karşıt inceleme / 3-4-3 safha', self.karsit_ve_safha),
                       ('3-4-4 Yüklenilen', self.yuklenilen),
                       ('3.7 Tevkifat', self.tevkifat),
                       ('3.8 Beyanname dökümü', self.dokum),
                       ('Metin paragrafları', self.paragraflar),
                       ('Son kontroller', self.son_kontroller)]:
            self.adim(ad, fn)
        cikti_yolu.parent.mkdir(parents=True, exist_ok=True)
        self.d.save(str(cikti_yolu))
        return Sonuc(cikti_yolu, self.kl, self.ozet)

    # ------------------------------------------------------------ yardımcılar
    def tablo(self, ad):
        return self.T.get(ad)

    def oran(self, x):
        return x / self.base if self.base else 0.0

    def _sablon_donemi(self):
        kap = self.tablo('kapak')
        eski = None
        if kap is not None:
            r = find_row(kap, r'İncelemenin Dönemi')
            if r is not None:
                m = re.search(r'([A-ZÇĞİÖŞÜ]+)-(\d{4})', uniq_cells(r)[-1].text)
                if m:
                    try:
                        eski = Donem.coz(f'{m.group(1)}-{m.group(2)}')
                    except ValueError:
                        pass
        if eski is None:
            raise SablonHatasi('Şablonun dönemi kapaktaki "İncelemenin Dönemi" satırından okunamadı.')
        self.eski = eski
        if eski == self.yeni:
            self.kl.hata('Şablon', f'Şablon zaten {eski.tire} raporu; şablon olarak bir önceki ayın raporu verilmeli.')
        elif eski.sonraki() != self.yeni:
            self.kl.uyari('Şablon', f'Şablon {eski.tire} raporu, girdiler {self.yeni.tire} dönemine ait; şablon bir önceki '
                          f'ayın ({self.yeni.onceki().tire}) raporu değil — dönem metinleri ve devreden kontrolü yanıltıcı olabilir.')
        # şablondaki "sonraki döneme devreden" (kontrol 10 için) — doldurmadan önce okunur
        self.onceki_sonraki = None
        for ad, pat in (('is_hacmi', r'^SONRAKİ'), ('dokum', r'^Sonraki Döneme')):
            t = self.tablo(ad)
            r = find_row(t, pat) if t is not None else None
            if r is not None and PARA_RE.match(uniq_cells(r)[-1].text.strip()):
                self.onceki_sonraki = num(uniq_cells(r)[-1].text.strip())
                break

    # ------------------------------------------------------------ 1) Ocak'a atıf yapan paragraflar
    def ocak_paragraflari(self):
        yil, ay = self.yeni.yil, self.yeni.ay
        ref = self.ayar.ocak_referansi(yil)
        if ay != 1 and not ref:
            ref = f'{DOLD} tarih ve {self.ayar.rapor_oneki}/{DOLD}'
            self.kl.elle('Ocak raporu atfı', f'OCAK-{yil} raporunun tarihi/sayısı ayar dosyasında yok (ocak_raporu veya '
                         f'rapor_referanslari "OCAK/{yil}"); 1-1, 1-3, 1-8, 1-9, 1-10 paragraflarında [DOLDURULACAK].')
        n = 0
        for p in self.d.paragraphs:
            t = p.text
            if 'raporumuza eklen' not in t:
                continue
            if ay == 1:
                deg = paragraf_degistir(p, r'\S+ tarih ve YMM \S+ sayılı OCAK-\d{4} dönemi raporumuza eklendiğinden bu raporumuza eklenmemiştir',
                                        f'OCAK-{yil} dönemi raporumuza eklenmiştir')
            else:
                deg = paragraf_degistir(p, r'(?:\S+ tarih ve YMM \S+ sayılı )?OCAK-\d{4} dönemi raporumuza (eklenmiştir|eklendiğinden bu raporumuza eklenmemiştir)',
                                        f'{ref} sayılı OCAK-{yil} dönemi raporumuza eklendiğinden bu raporumuza eklenmemiştir')
            n += deg
            if DOLD in p.text:
                isaretle(p)
        if ay == 1:
            self.kl.elle('Ocak raporu ekleri', 'Ocak raporu: ticaret sicil, kapasite raporu, gelir tablosu, imza sirküleri ve '
                         'nüfus cüzdanı suretleri bu rapora eklenmeli (1-1, 1-3, 1-8, 1-9, 1-10 "eklenmiştir" yazıldı).')
        self.ozet['ocak_paragraf'] = n

    # ------------------------------------------------------------ 2) dönem metinleri
    def donem_metinleri(self):
        e, y = self.eski, self.yeni
        reps = [(e.egik, y.egik), (e.tire, y.tire), (f'{e.yil} {e.ad} KDV', f'{y.yil} {y.ad} KDV')]
        for p in self.d.paragraphs:
            if 'OCAK-' in p.text and 'raporumuza' in p.text and y.ay != 1:
                continue
            for a, b in reps:
                if a in p.text:
                    metin_degistir(p, a, b)
        for yer, p in dx.tum_paragraflar(self.d):
            if yer in ('üst bilgi', 'alt bilgi'):
                for a, b in reps:
                    if a in p.text:
                        metin_degistir(p, a, b)

    # ------------------------------------------------------------ 3) teminat mektubu tarih / no
    def teminat_paragraflari(self):
        tem = self.g.tem
        pars = [p for p in self.d.paragraphs if 'Teminat Mektub' in p.text]
        if not tem.get('tarih'):
            for p in pars:
                isaretle(p)
            self.kl.hata('Teminat mektubu', 'Teminat mektubu kabul dilekçesi bulunamadı ya da tarih/numara okunamadı; '
                         f'{len(pars)} paragraftaki teminat tarihi/numarası önceki aydan kaldı (sarı) — elle girilmeli.', no='7')
            return
        n = 0
        for p in pars:
            n += paragraf_degistir(p, r'\d\d\.\d\d\.\d{4} tarih ve \d+ numaralı', f"{tem['tarih']} tarih ve {tem['no']} numaralı")
        guncel = [p for p in pars if f"{tem['tarih']} tarih ve {tem['no']} numaralı" in p.text]
        self.kl.sonuc(bool(pars) and len(guncel) == len(pars), 'Teminat mektubu',
                      f"Teminat mektubu tarihi/numarası {len(pars)} paragrafta {tem['tarih']} / {tem['no']} yapıldı.",
                      f"Teminat mektubu tarihi/numarası yalnızca {len(guncel)}/{len(pars)} paragrafta güncellenebildi "
                      '(kalıp "gg.aa.yyyy tarih ve NNN numaralı" değil) — elle kontrol edin.', no='7')
        for p in pars:
            if p not in guncel:
                isaretle(p)
        if tem.get('banka'):
            banka = ' '.join(katla(tem['banka']).upper().split())
            for p in pars:
                if banka not in ' '.join(katla(p.text).upper().split()):
                    isaretle(p)
                    self.kl.uyari('Teminat bankası', f'Dilekçedeki banka "{tem["banka"]}" paragrafta aynı yazımla geçmiyor; '
                                  f'banka/şube adını kontrol edin: "{p.text[:120]}…"', no='7')

    # ------------------------------------------------------------ 4) kapak
    def kapak(self):
        T0 = self.tablo('kapak')
        if T0 is None:
            return
        r = find_row(T0, r'^Rapor Sayısı')
        if r is not None:
            set_cell(uniq_cells(r)[1], f'{self.ayar.rapor_oneki}/{DOLD}', SARI)
        r = find_row(T0, r'^Rapor Ekleri')
        if r is not None:
            for c in uniq_cells(r)[2:]:
                set_cell(c, DOLD, SARI)
        r = find_row(T0, r'İncelemenin Dönemi')
        set_cell(uniq_cells(r)[-1], self.yeni.tire)
        for r in find_rows(T0, r'^Dayanak Sözleşmenin'):
            hucre_isaretle(uniq_cells(r)[-1])
        self.kl.elle('Kapak', 'Rapor sayısı ve tarihi [DOLDURULACAK]; dayanak sözleşme günü/sayısı önceki rapordan alındı '
                     '(sarı) — kontrol edin.')

    # ------------------------------------------------------------ 5) işçi tablosu (12 ay kayan)
    def isci(self):
        T6 = self.tablo('isci')
        if T6 is None:
            return
        if len(T6.columns) != 4:
            raise SablonYok(f'işçi tablosu 4 sütun değil ({len(T6.columns)})')
        cells = []
        for col in (0, 2):
            for r in T6.rows:
                cs = r.cells
                cells.append((cs[col], cs[col + 1]))
        vals = [(a.text.strip(), b.text.strip()) for a, b in cells]
        if vals[-1][0] != self.eski.tire:
            self.kl.uyari('İşçi sayısı', f'İşçi tablosunun son ayı {vals[-1][0]}, şablon dönemi {self.eski.tire} — kaydırma kontrol edilmeli.')
        vals = vals[1:] + [(self.yeni.tire, DOLD)]
        for (a, b), (x, y) in zip(cells, vals):
            set_cell(a, x)
            set_cell(b, y, SARI if y == DOLD else None)
        self.kl.elle('İşçi sayısı', f'1-6 işçi sayısı: {self.yeni.tire} işçi sayısı SGK bildirgesinden girilmeli.')

    # ------------------------------------------------------------ 6) matrah ve vergi bildirimi (T8 ve T15)
    def _bolum(self, T, baslangic, bitis):
        """(başlık satırı, sütun başlığı satırı|None, veri satırları) — bölüm yoksa None."""
        rows = list(T.rows)
        i = next((i for i, r in enumerate(rows) if ara(baslangic, row_text(r))), None)
        if i is None:
            return None
        kolon, veri = None, []
        for r in rows[i + 1:]:
            txt = row_text(r)
            if any(ara(b, txt) for b in bitis):
                break
            if ara(r'^\|?\s*(İşlem Türü|İstisna Türü|KDV Oranı)\b', txt):
                kolon = r
                continue
            veri.append(r)
        return rows[i], kolon, veri

    def _grup(self, T, tablo_adi, bolum_adi, bolum, istenen, kalem_anahtari, satir_anahtari, yaz, etiket, sablon=None):
        """Bölümdeki veri satırlarını istenen kalemlere eşitler ve yazar."""
        bas, kolon, veri = bolum
        taninmayan = [r for r in veri if satir_anahtari(r) is None]
        veri = [r for r in veri if satir_anahtari(r) is not None]
        for r in taninmayan:
            for c in uniq_cells(r):
                hucre_isaretle(c)
            self.kl.uyari('Tanınmayan satır', f'{tablo_adi} / {bolum_adi}: "{ilk_hucre(r).text.strip()[:80]}" satırı program '
                          'tarafından tanınmadı, güncellenmedi (sarı) — elle kontrol edin.')
        if not veri and sablon is None:
            sablon = kolon or bas
        sonuc, silinen = grubu_esitle(veri, istenen, kalem_anahtari, satir_anahtari, sonra=kolon or bas, sablon=sablon)
        self._silinenler(f'{tablo_adi} / {bolum_adi}', silinen)
        for row, kalem, yeni in sonuc:
            if yeni:
                set_cell(ilk_hucre(row), etiket(kalem), SARI)
                self.kl.bilgi('Yeni satır', f'{tablo_adi} / {bolum_adi}: "{etiket(kalem)}" satırı eklendi (sarı) — etiket ve biçimi kontrol edin.')
            try:
                yaz(row, kalem)
            except SablonYok as e:
                if self.kati:
                    raise
                for c in uniq_cells(row):
                    hucre_isaretle(c)
                self.kl.hata('Şablon yapısı', f'{tablo_adi} / {bolum_adi}: {e} — satır doldurulamadı (sarı), elle girin.')
        if not istenen and veri:
            self.kl.uyari('Boş bölüm', f'{tablo_adi} / {bolum_adi}: bu ay beyanda satır yok, şablondaki satırlar silindi; '
                          'bölüm başlığı duruyor — gerekiyorsa silin.')
        return sonuc

    def _silinenler(self, yer, silinen):
        for ad in silinen:
            self.kl.bilgi('Silinen satır', f'{yer}: "{ad[:90]}" satırı bu ay beyanda olmadığı için silindi.')

    def matrah_doldur(self, ad):
        T = self.tablo(ad)
        if T is None:
            return
        k = self.k
        baslik = 'Matrah tablosu' if ad == 'matrah' else 'Beyanname dökümü'
        # tevkifat uygulanmayan
        bol = self._bolum(T, r'TEVKİFAT UYGULANMAYAN', [r'KISMİ TEVKİFAT', r'DİĞER İŞLEMLER', r'^Matrah Toplamı'])
        if bol is None:
            raise SablonYok(f'{baslik}: "TEVKİFAT UYGULANMAYAN İŞLEMLER" bölümü yok')
        lines = []
        if k.get('r701'):
            lines.append(('701', [tr(k['r701'][0]), str(k['r701'][1]), tr(k['r701'][2])]))
        for a, b, c in k['yurtici']:
            lines.append(('yi', [tr(a), str(b), tr(c)]))
        self._grup(T, baslik, 'Tevkifat uygulanmayan', bol, lines, lambda x: x[0],
                   lambda r: '701' if '11/1-c' in ilk_hucre(r).text else ('yi' if ara(r'Yurtiçi Teslim', ilk_hucre(r).text) else None),
                   lambda r, x: set_row_vals(r, x[1]),
                   lambda x: 'İhracatı Yapılacak Nihai Ürünlerin Kanunun 11/1-c Maddesi Kapsamında Teslimi' if x[0] == '701' else 'Yurtiçi Teslim ve Hizmetler')
        # kısmi tevkifat
        tev = []
        if k.get('k410'):
            tev.append(('410', 'Yapım İşleri ile Bu İşlerle Birlikte İfa Edilen Müh.-Mim. Ve Etüt-Proje Hizmetleri [KDVGUT-(I/C-2.1.3.2.1)]', k['k410']))
        if k.get('k448'):
            tev.append(('448', 'Demir-Çelik Ürünlerinin Teslimi [KDVGUT-(I/C-2.1.3.3.8)]', k['k448']))
        bol = self._bolum(T, r'KISMİ TEVKİFAT UYGULANAN', [r'DİĞER İŞLEMLER', r'^Matrah Toplamı'])
        if bol is None:
            if tev:
                raise SablonYok(f'{baslik}: "KISMİ TEVKİFAT UYGULANAN İŞLEMLER" bölümü yok ama beyanda kısmi tevkifat var')
        else:
            self._grup(T, baslik, 'Kısmi tevkifat', bol, tev, lambda x: x[0],
                       lambda r: '410' if ara(r'Yapım İşleri', ilk_hucre(r).text) else ('448' if ara(r'Demir-?Çelik', ilk_hucre(r).text) else None),
                       lambda r, x: set_row_vals(r, [tr(x[2][0]), str(x[2][1]), x[2][2], tr(x[2][3])]),
                       lambda x: x[1])
        # diğer işlemler (amortisman satırı en sonda)
        dig = []
        if k.get('d_iade'):
            dig.append(('iade', 'Alınan Malların İadesi, Gerçekleşmeyen İşlemler', k['d_iade']))
        if k.get('d_kur'):
            dig.append(('kur', 'Kur Farkı / Yuvarlama Farkı Nedeniyle Oluşan KDV', k['d_kur']))
        if k.get('d_diger'):
            dig.append(('diger', 'Diğerleri', k['d_diger']))
        if k.get('d_amort'):
            dig.append(('amort', 'Amortismana Tabi Sabit Kıymet (Taşınmaz, Taşıt Araçları, Demirbaş, Makine ve Teçhizat vb.) Satışları', k['d_amort']))

        def dig_anahtar(r):
            t = ilk_hucre(r).text
            for a, pat in (('iade', r'Alınan Malların İadesi'), ('kur', r'Kur Farkı'), ('diger', r'^\s*Diğerleri'), ('amort', r'Amortismana Tabi')):
                if ara(pat, t):
                    return a
            return None
        bol = self._bolum(T, r'DİĞER İŞLEMLER', [r'^Matrah Toplamı'])
        if bol is None:
            if dig:
                raise SablonYok(f'{baslik}: "DİĞER İŞLEMLER" bölümü yok ama beyanda diğer işlemler var')
        else:
            self._grup(T, baslik, 'Diğer işlemler', bol, dig, lambda x: x[0], dig_anahtar,
                       lambda r, x: set_row_vals(r, [tr(x[2][0]), tr(x[2][1])]), lambda x: x[1])
        # toplamlar
        for pat, key in [(r'^Matrah Toplamı', 'matrah_toplam'), (r'^Hesaplanan KDV Toplamı', 'hesaplanan'),
                         (r'İlave', 'ilave'), (r'^Toplam KDV', 'toplam_kdv')]:
            r = find_row(T, pat)
            if r is not None:
                set_row_vals(r, [tr(k.get(key, 0))])

    def _oranlar(self, T, baslik):
        k = self.k
        orows = [r for r in T.rows if re.match(r'^%\d', row_text(r))]
        if not k['oranlar']:
            for r in orows:
                for c in uniq_cells(r):
                    hucre_isaretle(c)
            self.kl.hata('Oran dağılımı', f'{baslik}: KDV 1\'den oran dağılımı okunamadı; şablondaki satırlar bırakıldı (sarı) — elle girin.')
            return
        if not orows:
            raise SablonYok(f'{baslik}: oran dağılımı satırı (%1/%10/%20) yok')
        sonuc, silinen = grubu_esitle(orows, k['oranlar'], lambda x: x[0],
                                      lambda r: int(re.match(r'^%(\d+)', row_text(r)).group(1)))
        self._silinenler(f'{baslik} oran dağılımı', silinen)
        for r, (o, b, v), yeni in sonuc:
            set_cell(ilk_hucre(r), f'%{o}')
            set_row_vals(r, [tr(b), tr(v)])

    # ------------------------------------------------------------ 7) T9 indirimler + oranlar
    def indirimler(self):
        T9 = self.tablo('indirim')
        if T9 is None:
            return
        k = self.k
        set_row_vals(find_row(T9, r'^Yurtiçi Alımlara'), [tr(k['yurtici_alim'])])
        rs = find_row(T9, r'^Sorumlu')
        set_row_vals(rs, [tr(k.get('sorumlu', 0))])
        ri = find_row(T9, r'^İthalde')
        if k.get('ithal'):
            if ri is None:
                ri = dx.clone_row_after(rs)
                set_cell(ilk_hucre(ri), 'İthalde Ödenen KDV', SARI)
                self.kl.bilgi('Yeni satır', '3-2: "İthalde Ödenen KDV" satırı eklendi (sarı).')
            set_row_vals(ri, [tr(k['ithal'])])
        elif ri is not None:
            del_row(ri)
        self._oranlar(T9, '3-2')

    # ------------------------------------------------------------ 8) T10 iş hacmi
    def is_hacmi(self):
        T10 = self.tablo('is_hacmi')
        if T10 is None:
            return
        k = self.k
        for pat, val in [(r'^TESLİM VE HİZMET', k['aylik_bedel']), (r'^TOPLAM HES', k['toplam_kdv']),
                         (r'^ÖNCEKİ DÖNEM', k.get('devreden_onceki', 0)), (r'^YURTİÇİ ALIM', k['yurtici_alim']),
                         (r'^SORUMLU', k.get('sorumlu', 0)), (r'^İTHALDE', k.get('ithal', 0)),
                         (r'^SATIŞTAN', k.get('satis_iade', 0)), (r'^TOPLAM İND', k['indirim_toplam']),
                         (r'^İADE EDİLMESİ', k.get('iade_gereken', 0)), (r'^SONRAKİ', k.get('sonraki_devreden', 0))]:
            r = find_row(T10, pat)
            if r is not None:
                set_row_vals(r, [tr(val)])
        turler = [(kod, v[1]) for kod, v in sorted(self.k_kalem.items()) if v[1] > 0]
        mevcut = [r for r in T10.rows if re.match(r'^\d{3}\s*[–-]', row_text(r))]
        sonra = find_row(T10, r'^İADE EDİLMESİ')
        etiketler = dict(IADE_ETIKETLERI, **{str(a): b for a, b in self.ayar.iade_turu_etiketleri.items()})
        sonuc, silinen = grubu_esitle(mevcut, turler, lambda x: x[0], lambda r: re.match(r'^(\d{3})', row_text(r)).group(1),
                                      sonra=sonra, sablon=None if mevcut else sonra)
        self._silinenler('3-3-1 İş hacmi', silinen)
        for r, (kod, v), yeni in sonuc:
            if yeni:
                set_cell(ilk_hucre(r), etiketler.get(kod, kod), SARI)
                self.kl.bilgi('Yeni satır', f'3-3-1: "{etiketler.get(kod, kod)}" iade türü satırı eklendi (sarı) — etiket metnini kontrol edin.')
            set_row_vals(r, [tr(v)])
        kurallar.iade_turleri(self.kl, self.k_kalem, k.get('iade_gereken', 0), self.ayar.tolerans)
        self.ozet['iade_turleri'] = {kod: v for kod, v in turler}
        if self.ayar.iade_turleri:
            beklenmeyen = sorted({kod for kod, _ in turler} - set(self.ayar.iade_turleri))
            if beklenmeyen:
                self.kl.uyari('İade türü', f'Beyanda firma ayarında olmayan iade türü var: {", ".join(beklenmeyen)} '
                              f'(ayar: {", ".join(self.ayar.iade_turleri)}) — ilgili paragraf ve tabloları elle kontrol edin.')

    # ------------------------------------------------------------ 9) T11 karşıt + T12 safha
    def karsit_ve_safha(self):
        g = self.g
        base = self.base
        if base <= 0:
            self.kl.hata('Karşıt inceleme', 'Beyandaki indirilecek KDV (yurtiçi + sorumlu + ithal) sıfır; oranlar hesaplanamadı.')
        lst_total = sum(r['kdv'] for r in g.indirilecek)
        kurallar.liste_toplami(self.kl, lst_total, base)
        kurallar.liste_donemi(self.kl, g.liste_donemleri, self.yeni)
        s = safha_hesapla(g.takip, g.indirilecek, base, self.ayar, self.ayar.tolerans)
        _not_aktar(self.kl, s.notlar)
        self.safha = s
        self.ozet.update(liste_toplami=lst_total, karsit_toplam=s.tot, ekli_toplam=s.tot_e, muhafaza_toplam=s.tot_m,
                         ekli=[(r.firma, r.sekil, r.kdv) for r in s.ekli], muhafaza=[(r.firma, r.sekil, r.kdv) for r in s.muhafaza],
                         takip=[(r.firma, r.sekil, r.kdv, r.liste_kdv) for r in s.satirlar])
        if base > 0:
            self.kl.sonuc(s.tot >= base * ASGARI_ORAN - 0.005, 'Karşıt inceleme oranı',
                          f'Karşıt inceleme oranı {pct2(s.tot / base)} (≥ %80).',
                          f'Karşıt inceleme oranı {pct2(s.tot / base)} — %80 altında!', no='3')
            self.kl.sonuc(s.tot_e >= base * ASGARI_ORAN - 0.005, 'EKLİ toplamı',
                          f'EKLİ tutanakların oranı {pct2(s.tot_e / base)} (≥ %80).',
                          f'EKLİ tutanakların oranı {pct2(s.tot_e / base)} — %80 altında!', no='3')
        T11 = self.tablo('karsit')
        if T11 is not None:
            for pat, deger in ((r'İndirilecek KDV Tutarı', base), (r'x\s*%\s*80', base * ASGARI_ORAN)):
                r = find_row(T11, pat)
                if r is None:
                    raise SablonYok(f'karşıt inceleme tablosunda "{pat}" satırı yok')
                cs = uniq_cells(r)
                set_cell(cs[0], self.yeni.tire)
                set_cell(self._para_hucresi(cs, 2), tr(deger))
            r = find_row(T11, r'Karşıt İnceleme Yapılan Tutar')
            for x in uniq_cells(r):
                if re.fullmatch(r'[\d.]+,\d\d', x.text.strip()):
                    set_cell(x, tr(s.tot))
                if re.fullmatch(r'%[\d,]+', x.text.strip()):
                    set_cell(x, pct2(self.oran(s.tot)))
        T12 = self.tablo('safha')
        if T12 is not None:
            self._safha_tablosu(T12, s)
            belge.safha_tablosu_denetle(self.kl, T12)   # %80 oranı yukarıda ayrıca kontrol edildi

    @staticmethod
    def _para_hucresi(cs, varsayilan):
        for c in cs:
            if PARA_RE.match(c.text.strip()):
                return c
        return cs[varsayilan]

    def _safha_tablosu(self, T12, s):
        e, y = self.eski, self.yeni
        for r in T12.rows:
            c0 = uniq_cells(r)[0]
            for p in c0.paragraphs:
                if e.egik in p.text:
                    metin_degistir(p, e.egik, y.egik)
        rws = list(T12.rows)
        i_top = next((i for i, r in enumerate(rws) if i > 0 and 'TOPLAM' in row_text(r) and 'İNCELEME ORANI' in row_text(r)
                      and 'GENEL' not in row_text(r)), None)
        if i_top is None or 'GENEL' not in row_text(rws[-1]) or 'TOPLAM' not in row_text(rws[-2]) or i_top < 2:
            raise SablonYok('safha tablosunda EKLİ toplamı / muhafaza toplamı / GENEL TOPLAM satırları beklenen yerde değil')
        e_rows = rws[1:i_top]
        m_rows = rws[i_top + 1:-2]
        e_t = e_rows[0]
        if not s.ekli:
            for c in uniq_cells(e_t):
                hucre_isaretle(c)
            raise SablonYok('EKLİ olacak firma yok (takip listesinde KDV\'li firma bulunamadı); safha tablosu önceki aydan kaldı')
        for r in e_rows[1:]:
            del_row(r)
        cur = e_t
        for i, r in enumerate(s.ekli):
            rr = e_t if i == 0 else dx.clone_row_after(cur, e_t)
            cs = uniq_cells(rr)
            set_cell(cs[-5], r.firma)
            set_cell(cs[-4], r.sekil, SARI if DOLD in r.sekil else None)
            set_cell(cs[-3], f"{i + 1}- {'İTHALAT' if r.sekil == 'İTHALAT' else 'EKLİ'}")
            set_cell(cs[-2], tr(r.kdv))
            set_cell(cs[-1], pct(self.oran(r.kdv)))
            cur = rr
        tot_row = rws[i_top]
        set_cell(uniq_cells(tot_row)[-2], tr(s.tot_e))
        set_cell(uniq_cells(tot_row)[-1], pct(self.oran(s.tot_e)))
        # muhafaza satırları: % sütunu dikey birleşik (ilk satırda toplam oran)
        if not m_rows:
            if s.muhafaza:
                raise SablonYok('safha tablosunda "Muhafaza" satırı örneği yok')
        else:
            m_first = m_rows[0]
            if len(m_rows) > 1:
                cont_tpl = copy.deepcopy(m_rows[1]._tr)
            else:
                cont_tpl = copy.deepcopy(m_first._tr)
                sutun = dx.vmerge_baslangic_sutunu(m_first._tr)
                if sutun is not None:
                    dx.vmerge_devam_yap(cont_tpl, sutun)
            for r in m_rows[1:]:
                del_row(r)
            if not s.muhafaza:
                del_row(m_first)
                self.kl.uyari('Safha tablosu', 'Bu ay "Tarafımızdan Muhafaza Edilmektedir." satırı yok; muhafaza toplamı 0 yazıldı.')
            cur = m_first
            for i, r in enumerate(s.muhafaza):
                if i == 0:
                    rr = m_first
                else:
                    new = copy.deepcopy(cont_tpl)
                    cur._tr.addnext(new)
                    rr = _Row(new, T12)
                cs = uniq_cells(rr)
                idx = [j for j, cc in enumerate(cs) if 'Muhafaza' in cc.text]
                if not idx:
                    raise SablonYok('muhafaza satırında "Tarafımızdan Muhafaza Edilmektedir." hücresi yok')
                j = idx[0]
                set_cell(cs[j - 2], r.firma)
                set_cell(cs[j - 1], r.sekil, SARI if DOLD in r.sekil else None)
                set_cell(cs[j + 1], tr(r.kdv))
                if i == 0 and len(cs) > j + 2:
                    set_cell(cs[j + 2], pct(self.oran(s.tot_m)))
                cur = rr
        rws = list(T12.rows)
        set_cell(uniq_cells(rws[-2])[-2], tr(s.tot_m))
        set_cell(uniq_cells(rws[-2])[-1], pct(self.oran(s.tot_m)))
        set_cell(uniq_cells(rws[-1])[-2], tr(s.tot))
        set_cell(uniq_cells(rws[-1])[-1], pct(self.oran(s.tot)))

    # ------------------------------------------------------------ 10) T13 yüklenilen
    def yuklenilen(self):
        T13 = self.tablo('yuklenilen')
        k, yuk = self.k, self.g.yuklenilen
        hedef = k.get('301_yuklenilen', 0) + k.get('339_iade', 0)
        if not yuk:
            if k.get('301_yuklenilen'):
                self.kl.hata('Yüklenilen tutanak', f'Beyanda 301 yüklenilen KDV {tr(k["301_yuklenilen"])} var ama '
                             '"yüklenilen tutanak çalışması" dosyası bulunamadı/boş; 3-4-4 tablosu önceki aydan kaldı (sarı).', no='9')
                if T13 is not None:
                    for r in T13.rows:
                        for c in uniq_cells(r):
                            hucre_isaretle(c)
            elif T13 is not None:
                self.kl.uyari('Yüklenilen tutanak', 'Bu ay yüklenilen KDV yok ama şablonda 3-4-4 tablosu var — tabloyu/bölümü elle kaldırın.')
            return
        if T13 is None:
            self.kl.hata('Yüklenilen tutanak', '3-4-4 yüklenilen KDV tablosu şablonda bulunamadı; tutanaklar rapora yazılamadı.')
            return
        g = collections.OrderedDict()
        for x in yuk:
            g.setdefault((x['firma'], x['per']), 0.0)
            g[(x['firma'], x['per'])] += x['kdv']
        old_names = [uniq_cells(r)[1].text.strip() for r in T13.rows[1:-1]] + [t['firma'] for t in self.g.takip]

        def nice(n):
            for o in old_names:
                if firma_anahtari(o) == firma_anahtari(n):
                    return o
            return n
        ftot = collections.defaultdict(float)
        for (f, p), v in g.items():
            ftot[f] += v
        items = sorted(g.items(), key=lambda kv: (-ftot[kv[0][0]], kv[0][1]))
        ytot = sum(v for _, v in items)
        rws = list(T13.rows)
        if len(rws) < 3:
            raise SablonYok('yüklenilen tablosunda örnek satır yok')
        tmpl = rws[1]
        for r in rws[2:-1]:
            del_row(r)
        cur = tmpl
        refs = self.ayar.rapor_referanslari
        eksik_ref = []
        for i, ((f, (yy, mm)), v) in enumerate(items):
            rr = tmpl if i == 0 else dx.clone_row_after(cur, tmpl)
            cs = uniq_cells(rr)
            per = f'{AYLAR[mm - 1]}/{yy}'
            ref = refs.get(per) if (yy, mm) != (self.yeni.yil, self.yeni.ay) else None
            set_cell(cs[1], nice(f))
            set_cell(cs[2], 'BİLGİ İSTEME')
            set_cell(cs[3], per)
            set_cell(cs[4], f'{i + 1}- EKLİ')
            if ref:
                set_cell(cs[5], f'{ref[0]} – {self.ayar.rapor_oneki}/{ref[1]} Sayılı Raporda Mevcut')
            else:
                set_cell(cs[5], f'{DOLD} – {self.ayar.rapor_oneki}/{DOLD} Sayılı Raporda Mevcut', SARI)
                eksik_ref.append(per)
            set_cell(cs[6], tr(v))
            set_cell(cs[7], pct(v / ytot) if ytot else '')
            cur = rr
        last = list(T13.rows)[-1]
        set_cell(uniq_cells(last)[-2], tr(ytot))
        payda = k.get('301_yuklenilen') or ytot
        set_cell(uniq_cells(last)[-1], pct(ytot / payda) if payda else '')
        if k.get('301_yuklenilen'):
            kurallar.yuklenilen(self.kl, ytot, hedef, self.ayar.tolerans)
        if eksik_ref:
            self.kl.elle('Asıl tutanakların durumu', '3-4-4: şu dönemlerin rapor tarihi/sayısı ayar dosyasında (rapor_referanslari) '
                         f'yok: {", ".join(sorted(set(eksik_ref)))} — [DOLDURULACAK].')
        self.kl.bilgi('Yüklenilen inceleme şekli', '3-4-4: inceleme şekli tüm satırlarda "BİLGİ İSTEME" varsayıldı; tutanak türünü kontrol edin.')
        self.ozet['yuklenilen'] = [(f, f'{AYLAR[mm - 1]}/{yy}', v) for (f, (yy, mm)), v in items]

    # ------------------------------------------------------------ 11) T14 tevkifat
    def tevkifat(self):
        T14 = self.tablo('tevkifat')
        k = self.k
        if not k.get('k410'):
            if T14 is not None:
                for c in uniq_cells(T14.rows[-1]):
                    hucre_isaretle(c)
                self.kl.uyari('3.7 Tevkifat', 'Bu ay beyanda 410 kısmi tevkifat yok ama şablonda 3.7 tevkifat tablosu var '
                              '(önceki ay rakamları, sarı) — bölümü elle kaldırın.')
            return
        if T14 is None:
            self.kl.hata('3.7 Tevkifat', 'Beyanda 410 kısmi tevkifat var ama şablonda 3.7 tevkifat tablosu bulunamadı.')
            return
        matrah = k.get('410_tutar', k['k410'][0])
        oran = int(k['k410'][1])
        tev = k['k410'][2]
        hes = matrah * oran / 100
        iade = k.get('410_iade', 0)
        set_row_vals(T14.rows[-1], [self.yeni.tire, tr(matrah), f'%{oran}', tr(hes), tev, tr(iade)], start=0)
        kurallar.tevkifat_410(self.kl, matrah, oran, tev, iade)

    # ------------------------------------------------------------ 12) T15 beyanname dökümü
    def dokum(self):
        T15 = self.tablo('dokum')
        if T15 is None:
            return
        k = self.k
        self.matrah_doldur('dokum')
        for pat, val in [(r'^Satıştan İade', k.get('satis_iade', 0)), (r'^Yurtiçi Alımlarına', k['yurtici_alim']),
                         (r'^Sorumlu', k.get('sorumlu', 0)), (r'^İthalde', k.get('ithal', 0)),
                         (r'^Önceki Dönemden Devreden KDV', k.get('devreden_onceki', 0)), (r'^İndirimler Toplamı', k['indirim_toplam']),
                         (r'^İhraç Kaydıyla Teslim Bedeli Toplamı', k.get('ihrac_kayitli_bedel', 0)),
                         (r'^Tecil Edilebilir', k.get('ihrac_kayitli_iade', 0)),
                         (r'^İhracatın Gerçekleştiği Dönemde İade Edilecek Tecil', k.get('ihrac_kayitli_iade', 0)),
                         (r'^İhracatın Gerçekleştiği Dönemde İade Edilecek KDV', k.get('ihrac_kayitli_iade', 0)),
                         (r'^Tecil Edilecek', k.get('tecil_edilecek', 0)), (r'^Ödenmesi Gereken', k.get('odenmesi_gereken', 0)),
                         (r'^İade Edilmesi Gereken', k.get('iade_gereken', 0)), (r'^Sonraki Döneme', k.get('sonraki_devreden', 0)),
                         (r'\(Aylık\)', k['aylik_bedel']), (r'\(Kümülâtif\)', k.get('kumulatif') or 0)]:
            r = find_row(T15, pat)
            if r is not None:
                set_row_vals(r, [tr(val)])
        r = find_row(T15, r'^MATRAH VE VERGİ BİLDİRİMİ')
        if r is not None:
            set_cell(uniq_cells(r)[-1], f'MATRAH VE VERGİ BİLDİRİMİ-{self.yeni.egik}')
        self._oranlar(T15, '3.8')
        # ihraç kayıtlı bölümündeki 11/1-c satırı
        bol = self._bolum(T15, r'İHRAÇ KAYDIYLA TESLİMLERE', [r'^İhraç Kaydıyla Teslim Bedeli Toplamı'])
        if bol is not None:
            for r in bol[2]:
                if '11/1-c' not in ilk_hucre(r).text:
                    continue
                if k.get('r701'):
                    set_row_vals(r, [tr(k['r701'][0]), str(k['r701'][1]), tr(k['r701'][2])])
                else:
                    set_row_vals(r, [tr(0), '', tr(0)])
                    for c in uniq_cells(r):
                        hucre_isaretle(c)
                    self.kl.uyari('3.8 İhraç kayıtlı', 'Bu ay beyanda 701 (ihraç kayıtlı teslim) yok ama dökümde ihraç kaydıyla '
                                  'teslimler bölümü var (sarı) — bölümü elle kaldırın.')
        # tam istisna (3xx) ve kısmi tevkifat (4xx) iade satırları
        tam_etk = {'301': '301-Mal İhracatı', '318': '318-3996 Sayılı Kanuna Göre Yap İşlet Devret Modeli Çer.',
                   '339': '339-İmalat Sanayii ile Turizme Yönelik Yatırım Teşvik Belgesi Kapsamında Yapılan Teslim ve Hizmetler',
                   '350': '350-Digerleri'}
        kis_etk = {'410': '410-Yapım İşleri ile Bu İşlerle Birlikte İfa', '448': '448-Demir Çelik Ürünlerinin Teslimi'}
        r3 = [r for r in T15.rows if re.match(r'^3\d\d\s*-', row_text(r))]
        r4 = [r for r in T15.rows if re.match(r'^4\d\d\s*-', row_text(r))]
        for etk, mevcut, diger, bas in ((tam_etk, r3, r4, r'TAM İSTİSNA KAPSAMINA'), (kis_etk, r4, r3, r'KISMİ TEVKİFAT KAPSAMINA'),):
            istenen = [(kod, etk[kod], *self.k_kalem[kod]) for kod in IADE_KODLARI if kod in etk and kod in self.k_kalem]
            bas_r = find_row(T15, bas)
            sonra = None
            if bas_r is not None:
                rows = list(T15.rows)
                i = next(j for j, x in enumerate(rows) if x._tr is bas_r._tr)
                sonra = rows[i + 1] if i + 1 < len(rows) and ara(r'Türü', row_text(rows[i + 1])) else bas_r
            if not mevcut and not istenen:
                continue
            if not mevcut and sonra is None:
                raise SablonYok(f'3.8 dökümünde "{bas}" bölümü yok')
            sonuc, silinen = grubu_esitle(mevcut, istenen, lambda x: x[0], lambda r: re.match(r'^(\d{3})', row_text(r)).group(1),
                                          sonra=sonra, sablon=None if mevcut else (diger[0] if diger else sonra))
            self._silinenler('3.8 Beyanname dökümü', silinen)
            for r, (kod, lab, a, b), yeni in sonuc:
                set_cell(ilk_hucre(r), lab, SARI if yeni else None)
                set_row_vals(r, [tr(a), tr(b)])
                if yeni:
                    self.kl.bilgi('Yeni satır', f'3.8 dökümü: "{lab}" satırı eklendi (sarı).')

    # ------------------------------------------------------------ 13) metin paragrafları
    def paragraflar(self):
        k, tem = self.k, self.g.tem
        kalan = k.get('iade_gereken', 0) - tem.get('toplam', 0)
        kurallar_ = [
            ('3.5 (701)', lambda t: t.lstrip().startswith('Firmanın') and '11/1-c maddesi kapsamında toplam' in t,
             [(r'toplam [\d.]+,\d\d TL teslimde', f"toplam {tr(k.get('ihrac_kayitli_bedel', 0))} TL teslimde"),
              (r'toplam [\d.]+,\d\d TL KDV iadesi', f"toplam {tr(k.get('ihrac_kayitli_iade', 0))} TL KDV iadesi")],
             bool(k.get('r701') or k.get('ihrac_kayitli_bedel'))),
            ('3-6 (301)', lambda t: 'ihracatından dolayı toplam' in t,
             [(r'toplam [\d.]+,\d\d- TL teslimlerde', f"toplam {tr(k.get('301_tutar', 0))}- TL teslimlerde"),
              (r'toplam [\d.]+,\d\d- TL KDV iadesi', f"toplam {tr(k.get('301_yuklenilen', 0))}- TL KDV iadesi")],
             '301' in self.k_kalem),
            ('3.7 (410)', lambda t: 'YAPIM İŞLERİ (410)' in t,
             [(r'tutar [\d.]+,\d\d TL', f"tutar {tr(k.get('410_iade', 0))} TL")], bool(k.get('410_iade'))),
            ('3-8-1 Düzeltme', lambda t: re.match(r'\s*Firmanın .* KDV beyannamesine düzeltme', t),
             [(r'düzeltme (verilmiştir|verilmemiştir)', 'düzeltme verilmiştir' if k.get('duzeltme') else 'düzeltme verilmemiştir')], True),
            ('4-3', lambda t: t.startswith('4-3.'),
             [(r'Tesliminden dolayı [\d.]+,\d\d TL', f"Tesliminden dolayı {tr(k.get('ihrac_kayitli_iade', 0))} TL")],
             bool(k.get('r701') or k.get('ihrac_kayitli_iade'))),
            ('4-4', lambda t: t.startswith('4-4.'),
             [(r'301-Mal İhracatından dolayı [\d.]+,\d\d TL', f"301-Mal İhracatından dolayı {tr(k.get('301_yuklenilen', 0))} TL"),
              (r'Tutulan dolayı [\d.]+,\d\d TL', f"Tutulan dolayı {tr(k.get('410_iade', 0))} TL")], True),
            ('4-5', lambda t: t.startswith('4-5.'),
             ([(r'301-Mal İhracatından dolayı [\d.]+,\d\d TL', f"301-Mal İhracatından dolayı {tr(tem['301'])} TL")] if '301' in tem else [])
             + ([(r'Tevkifata Tabi Tutulan [\d.]+,\d\d TL', f"Tevkifata Tabi Tutulan {tr(tem['410'])} TL")] if '410' in tem else [])
             + ([(r'toplam [\d.]+,\d\d TL iade', f"toplam {tr(tem['toplam'])} TL iade")] if 'toplam' in tem else []), True),
            ('4-6', lambda t: t.startswith('4-6.'), [(r'KDV’nin [\d.]+,\d\d- TL', f'KDV’nin {tr(kalan)}- TL')], True),
            ('4-8', lambda t: t.startswith('4-8.'),
             [(r'tutarının [\d.]+,\d\d--TL', f"tutarının {tr(k.get('sonraki_devreden', 0))}--TL")], True),
        ]
        bulunan = collections.defaultdict(list)
        for p in self.d.paragraphs:
            t = p.text
            for ad, eslesir, degisimler, gecerli in kurallar_:
                if not eslesir(t):
                    continue
                bulunan[ad].append(p)
                tutmayan = [d for d, _ in degisimler if not re.search(d, p.text)]
                for d, yeni in degisimler:
                    paragraf_degistir(p, d, yeni)
                if tutmayan:
                    isaretle(p)
                    self.kl.hata('Paragraf güncellenemedi', f'{ad}: kalıp tutmadığı için tutar güncellenemedi (sarı) — '
                                 f'elle düzeltin: "{t[:150]}…"', yer=f'Paragraf {ad}')
                if not gecerli:
                    isaretle(p)
                    self.kl.uyari('Geçersiz paragraf', f'{ad}: bu ay beyanda ilgili işlem yok ama paragraf duruyor (sarı) — '
                                  'paragrafı/bölümü elle kaldırın.', yer=f'Paragraf {ad}')
                break
        for ad, _, _, gecerli in kurallar_:
            if gecerli and not bulunan[ad] and ad not in ('3-8-1 Düzeltme',):
                self.kl.uyari('Paragraf bulunamadı', f'{ad} paragrafı şablonda bulunamadı; tutarları elle kontrol edin.', yer=f'Paragraf {ad}')
        # 4-4 / 4-5 paragraflarında geçen iade türleri beyan/dilekçe ile aynı mı
        beyandaki = {kod for kod, (_, v) in self.k_kalem.items() if v > 0}
        dilekcedeki = {kod for kod in ('301', '318', '339', '350', '410', '448') if kod in tem}
        for ad, beklenen in (('4-4', beyandaki), ('4-5', dilekcedeki if tem else None)):
            if beklenen is None:
                continue
            for p in bulunan[ad]:
                gecen = set(re.findall(r'\b(3\d\d|4\d\d)-', p.text)) & set(IADE_KODLARI)
                if gecen != beklenen:
                    isaretle(p)
                    self.kl.uyari('İade türleri paragrafı', f'{ad} paragrafında geçen türler {sorted(gecen)}, '
                                  f'{"beyanda" if ad == "4-4" else "teminat dilekçesinde"} {sorted(beklenen)} — paragrafı elle düzenleyin (sarı).',
                                  yer=f'Paragraf {ad}')
        for ad, kodlar in (('4-4', beyandaki), ('4-5', dilekcedeki)):
            desteksiz = sorted(kodlar - {'301', '410'})
            if desteksiz and bulunan[ad]:
                for p in bulunan[ad]:
                    isaretle(p)
                self.kl.hata('Paragraf güncellenemedi', f'{ad}: {", ".join(desteksiz)} türü tutarları otomatik güncellenmiyor '
                             '(yalnızca 301 ve 410) — elle yazın (sarı).', yer=f'Paragraf {ad}')
        if 'toplam' not in tem:
            for p in bulunan['4-6'] + bulunan['4-5']:
                isaretle(p)
        self.ozet['teminatla_alinamayan'] = kalan

    # ------------------------------------------------------------ 14) son kontroller
    def son_kontroller(self):
        k = self.k
        kurallar.beyan_tutarliligi(self.kl, k)
        kurallar.teminat(self.kl, self.g.tem, self.k_kalem, k.get('iade_gereken', 0), self.ayar.tolerans)
        kurallar.devreden(self.kl, self.onceki_sonraki, k.get('devreden_onceki', 0), self.ayar.tolerans)
        # parasal hadler paragrafı ve yıl ifadeleri
        y = self.yeni.yil
        for p in self.d.paragraphs:
            t = p.text
            if 'VUK. 353' in t or '2023 Yılı İçin' in t:
                self.kl.uyari('Parasal hadler', '3-4-2 parasal hadler paragrafı hâlâ "2023 Yılı – VUK 353 – 150.000/450.000" diyor '
                              '(diğer raporlarda 2025: 65.000/195.000). Güncel tutarları ofis teyit etsin (sarı).', no='11')
                isaretle(p)
            for pat, beklenen, ac in ((r'bir önceki yıl \((\d{4})', y - 1, 'bir önceki yıl'),
                                      (r'(\d{4}) yılına ait kanuni defter', y, 'defter tasdik yılı')):
                for m in re.finditer(pat, t):
                    if int(m.group(1)) != beklenen:
                        isaretle(p)
                        self.kl.uyari('Yıl ifadesi', f'"{ac}" {m.group(1)} yazıyor, {beklenen} olmalı (sarı): "{t[:120]}…"', no='11')
        # eski dönem adı taraması
        e = self.eski
        desenler = [(re.escape(e.egik), f'Önceki dönem ({e.egik}) kaldı'), (re.escape(e.tire), f'Önceki dönem ({e.tire}) kaldı'),
                    (re.escape(f'{e.yil} {e.ad} KDV'), f'Önceki dönem ({e.yil} {e.ad} KDV) kaldı')]
        indeks = {t._tbl: i for i, t in enumerate(self.d.tables)}
        haric = {f'tablo{indeks[self.T[a]._tbl]}' for a in ('isci', 'yuklenilen') if a in self.T and self.T[a]._tbl in indeks}
        atla = (lambda t: 'raporumuza' in t and 'OCAK-' in t) if e.ay == 1 else None
        belge.eski_metin_tara(self.kl, self.d, desenler, haric, atla)
        # doldurulan tablolarda güncellenmemiş tutar hücreleri
        adlar = {'matrah': '3-1 Matrah', 'indirim': '3-2 İndirimler', 'is_hacmi': '3-3-1 İş hacmi', 'karsit': '3-4 Karşıt',
                 'safha': '3-4-3 Safha', 'yuklenilen': '3-4-4 Yüklenilen', 'tevkifat': '3.7 Tevkifat', 'dokum': '3.8 Döküm'}
        belge.guncellenmemis_tutarlar(self.kl, {a: (b, self.T[a]) for a, b in adlar.items() if a in self.T})


def taslak_olustur(sablon, girdiler, ayar, cikti, kati=False):
    return TaslakOlusturucu(sablon, girdiler, ayar, kati).olustur(cikti)
