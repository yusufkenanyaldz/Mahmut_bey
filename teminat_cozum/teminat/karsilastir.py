"""Taslak ↔ ofisin bitmiş raporu karşılaştırması (öğrenme döngüsü).

Paragraf ve tablo farklarını listeler; ayrıca her bölümdeki tutarları (1.234,56 biçimli) ayrı sayar:
"tutar farkı" = gerçek raporda olup taslakta aynı bölümde bulunmayan tutar sayısı. Hedef: 0.
"""
import collections
import difflib
from dataclasses import dataclass, field

from docx import Document

from .ortak import para_bul
from .rapor.tablolar import tablolari_bul

KAPSAM_DISI = ('---', '+++', '@@')


def table_rows(t):
    out = []
    for r in t.rows:
        cs = []
        for c in r.cells:
            x = ' '.join(c.text.split())
            if not cs or cs[-1] != x:
                cs.append(x)
        out.append(' | '.join(cs))
    return out


def _fark(a, b):
    return [x for x in difflib.unified_diff(a, b, lineterm='', n=0) if not x.startswith(KAPSAM_DISI)]


def _tutar_farki(a_satirlar, b_satirlar):
    a = collections.Counter(t for s in a_satirlar for t in para_bul(s))
    b = collections.Counter(t for s in b_satirlar for t in para_bul(s))
    return sorted((a - b).elements()), sorted((b - a).elements())


@dataclass
class Bolum:
    ad: str
    farklar: list
    satir: int
    eksik_tutar: list = field(default_factory=list)   # gerçek raporda var, taslakta yok
    fazla_tutar: list = field(default_factory=list)   # taslakta var, gerçek raporda yok


@dataclass
class Karsilastirma:
    bolumler: list
    gercek_tablo: int
    taslak_tablo: int

    @property
    def tutar_farki(self):
        return sum(len(b.eksik_tutar) for b in self.bolumler)

    @property
    def farkli_satir(self):
        return sum(len([x for x in b.farklar if x.startswith('-')]) for b in self.bolumler)

    def metin(self):
        out = []
        for b in self.bolumler:
            if not b.farklar and not b.eksik_tutar:
                continue
            out.append(f'=== {b.ad}')
            out.extend(b.farklar)
            if b.eksik_tutar:
                out.append(f'  ! Gerçek raporda olup taslakta olmayan tutarlar ({len(b.eksik_tutar)}): {", ".join(b.eksik_tutar)}')
            if b.fazla_tutar:
                out.append(f'  ! Taslakta olup gerçek raporda olmayan tutarlar ({len(b.fazla_tutar)}): {", ".join(b.fazla_tutar)}')
            out.append('')
        out.append(f'Özet: gerçek {self.gercek_tablo} tablo, taslak {self.taslak_tablo} tablo; '
                   f'{self.farkli_satir} satır farklı; tutar farkı: {self.tutar_farki}.  (- gerçek rapor, + taslak)')
        return '\n'.join(out)


def _tablo_eslestir(R, G):
    """Tabloları önce adla (tablolar.py tanımları), kalanları sırayla eşleştirir."""
    rt, _ = tablolari_bul(R)
    gt, _ = tablolari_bul(G)
    ciftler = []
    r_kul, g_kul = set(), set()
    ri = {t._tbl: i for i, t in enumerate(R.tables)}
    gi = {t._tbl: i for i, t in enumerate(G.tables)}
    for ad in rt:
        if ad in gt:
            ciftler.append((ad, rt[ad], gt[ad]))
            r_kul.add(ri.get(rt[ad]._tbl))
            g_kul.add(gi.get(gt[ad]._tbl))
    kalan_r = [(i, t) for i, t in enumerate(R.tables) if i not in r_kul]
    kalan_g = [(i, t) for i, t in enumerate(G.tables) if i not in g_kul]
    for (i, a), (j, b) in zip(kalan_r, kalan_g):
        ciftler.append((f'Tablo {i}' + (f' (taslak {j})' if i != j else ''), a, b))
    for i, a in kalan_r[len(kalan_g):]:
        ciftler.append((f'Tablo {i} (taslakta yok)', a, None))
    for j, b in kalan_g[len(kalan_r):]:
        ciftler.append((f'Taslak tablo {j} (gerçekte yok)', None, b))
    return ciftler


def karsilastir(gercek_yol, taslak_yol):
    R, G = Document(str(gercek_yol)), Document(str(taslak_yol))
    a = [' '.join(p.text.split()) for p in R.paragraphs if p.text.strip()]
    b = [' '.join(p.text.split()) for p in G.paragraphs if p.text.strip()]
    eksik, fazla = _tutar_farki(a, b)
    bolumler = [Bolum('Paragraflar', _fark(a, b), len(a), eksik, fazla)]
    for ad, ta, tb in _tablo_eslestir(R, G):
        ra = table_rows(ta) if ta is not None else []
        gb = table_rows(tb) if tb is not None else []
        eksik, fazla = _tutar_farki(ra, gb)
        bolumler.append(Bolum(ad, _fark(ra, gb), len(ra), eksik, fazla))
    return Karsilastirma(bolumler, len(R.tables), len(G.tables))
