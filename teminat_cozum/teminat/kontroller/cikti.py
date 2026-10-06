"""Kontrol listesini Excel (.xlsx) ve HTML olarak yazar."""
import datetime
import html
from pathlib import Path

from ..ortak import tr
from .liste import DURUM_ACIKLAMA, DURUMLAR

RENK = {'HATA': 'F8D7DA', 'UYARI': 'FFF3CD', 'ELLE': 'FFF9C4', 'BİLGİ': 'E2E8F0', 'TAMAM': 'D1E7DD'}
RENK_YAZI = {'HATA': '842029', 'UYARI': '664D03', 'ELLE': '5C4B00', 'BİLGİ': '334155', 'TAMAM': '0F5132'}


def _ozet_satirlari(ozet):
    s = []
    def ekle(ad, v):
        if v is None:
            return
        s.append((ad, tr(v) if isinstance(v, float) else str(v)))
    ekle('Dönem', ozet.get('donem'))
    ekle('Şablon (önceki rapor)', ozet.get('sablon'))
    ekle('Şablonun dönemi', ozet.get('sablon_donemi'))
    ekle('İndirilecek KDV (beyan: yurtiçi + sorumlu + ithal)', ozet.get('indirilecek_beyan'))
    ekle('İndirilecek KDV listesi toplamı', ozet.get('liste_toplami'))
    ekle('Karşıt inceleme yapılan tutar (safha genel toplamı)', ozet.get('karsit_toplam'))
    ekle('EKLİ + İTHALAT toplamı', ozet.get('ekli_toplam'))
    ekle('Muhafaza toplamı', ozet.get('muhafaza_toplam'))
    b = ozet.get('indirilecek_beyan')
    if b:
        if ozet.get('karsit_toplam') is not None:
            s.append(('Karşıt inceleme oranı', f"%{tr(ozet['karsit_toplam'] / b * 100)}"))
        if ozet.get('ekli_toplam') is not None:
            s.append(('EKLİ oranı', f"%{tr(ozet['ekli_toplam'] / b * 100)}"))
    for kod, v in (ozet.get('iade_turleri') or {}).items():
        ekle(f'İade: {kod}', v)
    ekle('Teminatla alınamayan (4-6)', ozet.get('teminatla_alinamayan'))
    for ad, yol in (ozet.get('dosyalar') or {}).items():
        ekle(f'Girdi: {ad}', yol)
    return s


def xlsx_yaz(kl, ozet, yol):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = 'Kontrol Listesi'
    bas = ['Durum', 'Konu', 'Açıklama', 'Yer', 'Beklenen', 'Bulunan', 'Talimat md.']
    ws.append(bas)
    for c in ws[1]:
        c.font = Font(bold=True)
    for k in kl.sirali():
        ws.append([k.durum, k.konu, k.aciklama, k.yer, k.beklenen, k.bulunan, k.no])
        r = ws.max_row
        ws.cell(r, 1).fill = PatternFill('solid', fgColor=RENK[k.durum])
        ws.cell(r, 1).font = Font(bold=True, color=RENK_YAZI[k.durum])
        ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical='top')
        for col in (1, 2, 4, 5, 6, 7):
            ws.cell(r, col).alignment = Alignment(vertical='top')
    for col, w in zip('ABCDEFG', (9, 26, 100, 16, 16, 16, 10)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:G{ws.max_row}'

    oz = wb.create_sheet('Özet')
    oz.append(['Bilgi', 'Değer'])
    for c in oz[1]:
        c.font = Font(bold=True)
    for a, b in _ozet_satirlari(ozet):
        oz.append([a, b])
    oz.append([])
    oz.append(['Durum', 'Adet', 'Anlamı'])
    for d in DURUMLAR:
        oz.append([d, kl.say(d), DURUM_ACIKLAMA[d]])
    oz.column_dimensions['A'].width = 50
    oz.column_dimensions['B'].width = 40
    oz.column_dimensions['C'].width = 60

    if ozet.get('takip'):
        sf = wb.create_sheet('Safha (3-4-3)')
        sf.append(['Rapor durumu', 'Firma', 'İnceleme şekli', 'Takip KDV', 'İndirilecek liste KDV', 'Fark'])
        for c in sf[1]:
            c.font = Font(bold=True)
        ekli = {(f, v) for f, _, v in ozet.get('ekli', [])}
        sira = {f: i + 1 for i, (f, _, _) in enumerate(ozet.get('ekli', []))}
        for f, sek, v, lv in ozet['takip']:
            durum = (f"{sira[f]}- {'İTHALAT' if sek == 'İTHALAT' else 'EKLİ'}") if (f, v) in ekli else 'Muhafaza'
            sf.append([durum, f, sek, v, lv, (v - lv) if lv is not None else None])
            for col in (4, 5, 6):
                sf.cell(sf.max_row, col).number_format = '#,##0.00'
        for col, w in zip('ABCDEF', (14, 60, 16, 18, 20, 14)):
            sf.column_dimensions[col].width = w
    wb.save(str(yol))


def html_yaz(kl, ozet, yol, baslik='Kontrol Listesi'):
    e = html.escape
    satirlar = []
    for k in kl.sirali():
        satirlar.append(
            f'<tr class="{e(k.durum)}"><td><span class="rozet r-{e(k.durum)}">{e(k.durum)}</span></td><td>{e(k.konu)}</td>'
            f'<td>{e(k.aciklama)}</td><td>{e(k.yer)}</td><td class="sayi">{e(k.beklenen)}</td>'
            f'<td class="sayi">{e(k.bulunan)}</td><td>{e(k.no)}</td></tr>')
    ozet_html = ''.join(f'<tr><th>{e(a)}</th><td>{e(b)}</td></tr>' for a, b in _ozet_satirlari(ozet))
    sayac = ''.join(f'<span class="rozet r-{d}">{d}: {kl.say(d)}</span> ' for d in DURUMLAR)
    filtre = ''.join(f'<label><input type="checkbox" checked data-d="{d}"> {d}</label> ' for d in DURUMLAR)
    stil = ''.join(f'.r-{d}{{background:#{RENK[d]};color:#{RENK_YAZI[d]}}}' for d in DURUMLAR)
    zaman = datetime.datetime.now().strftime('%d.%m.%Y %H:%M')
    icerik = f'''<!doctype html>
<html lang="tr"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{e(baslik)}</title>
<style>
body{{font-family:Segoe UI,Arial,sans-serif;margin:16px;color:#1f2937;background:#fff}}
h1{{font-size:20px;margin:0 0 4px}} .alt{{color:#6b7280;font-size:13px;margin-bottom:12px}}
table{{border-collapse:collapse;width:100%;font-size:13px;margin-bottom:20px}}
th,td{{border:1px solid #d1d5db;padding:5px 7px;vertical-align:top;text-align:left}}
thead th{{background:#f3f4f6;position:sticky;top:0}} td.sayi{{white-space:nowrap;text-align:right}}
.rozet{{display:inline-block;border-radius:4px;padding:1px 6px;font-weight:600;font-size:12px}}
{stil} .filtre{{margin:8px 0}} .ozet th{{width:45%;background:#f9fafb}}
@media print{{.filtre{{display:none}}}}
</style></head><body>
<h1>{e(baslik)}</h1><div class="alt">{e(str(ozet.get('donem', '')))} — oluşturma: {zaman}. Taslak programın ürettiği kontrol listesidir; karar ve imza Yeminli Mali Müşavire aittir.</div>
<div>{sayac}</div>
<div class="filtre">Göster: {filtre}</div>
<table><thead><tr><th>Durum</th><th>Konu</th><th>Açıklama</th><th>Yer</th><th>Beklenen</th><th>Bulunan</th><th>Md.</th></tr></thead>
<tbody>{''.join(satirlar)}</tbody></table>
<h2 style="font-size:16px">Özet</h2><table class="ozet">{ozet_html}</table>
<script>
document.querySelectorAll('.filtre input').forEach(function(c){{c.addEventListener('change',function(){{
document.querySelectorAll('tbody tr.'+CSS.escape(c.dataset.d)).forEach(function(r){{r.style.display=c.checked?'':'none'}})}})}});
</script></body></html>'''
    Path(yol).write_text(icerik, encoding='utf-8')


def konsol_ozeti(kl):
    satir = [' · '.join(f'{d}: {kl.say(d)}' for d in DURUMLAR)]
    for k in kl.sirali():
        if k.durum in ('HATA', 'UYARI'):
            satir.append(f'  [{k.durum}] {k.konu}: {k.aciklama}')
    return '\n'.join(satir)
