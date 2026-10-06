"""Testler için sentetik (uydurma) girdi klasörü ve şablon rapor üretir.

Firma adları ve tutarlar tamamen uydurmadır; gerçek mükellef verisi içermez. KDV 1 beyannamesi PDF yerine
pdfplumber'ın ürettiği metin biçiminde üretilir (testlerde `pdf_metni` bu metni döndürecek şekilde yamalanır).
"""
import copy
import datetime
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document

from teminat.ortak import Donem, tr

ITH = '1111111111'


@dataclass
class Tedarikci:
    firma: str
    vkn: str
    kdv: float
    ymm: str = ''
    fatura: int = 2
    takipte: bool = True


def varsayilan_tedarikciler():
    t = [
        Tedarikci('ALFA ÇELİK SAN. VE TİC. A.Ş.', '0010000001', 5_000_000.00, 'YMM ALİ VELİ'),
        Tedarikci('BETA YAPI MALZ. İNŞ. LTD. ŞTİ.', '0010000002', 3_200_000.00),
        Tedarikci('GAMA DEMİR SAN. A.Ş.', '0010000003', 2_100_000.00, 'YMM AYŞE FATMA'),
        Tedarikci('DELTA GALVANİZ YAPI ELEM. A.Ş.', '0010000004', 1_500_000.00, 'YMM AYŞE FATMA'),
        Tedarikci('EPSİLON METAL A.Ş.', '0010000005', 1_250_000.00, 'YMM ALİ VELİ'),
        Tedarikci('ZETA NAKLİYE', '12345678901', 1_000_000.00),
        Tedarikci('ETA PROFİL SAN. LTD. ŞTİ.', '0010000007', 900_000.00),
        Tedarikci('TETA ALÜMİNYUM A.Ş.', '0010000008', 800_000.00, 'YMM CAN CAN'),
        Tedarikci('İOTA ÇELİK SAN. A.Ş.', '0010000009', 700_000.00, 'YMM CAN CAN'),
        Tedarikci('KAPPA ORMAN ÜRÜN. LTD. ŞTİ.', '0010000010', 650_000.00, 'YMM CAN CAN'),
        Tedarikci('LAMBDA YALITIM A.Ş.', '0010000011', 640_000.00, 'YMM CAN CAN'),
        Tedarikci('MU İNŞAAT A.Ş.', '0010000012', 400_000.00, 'YMM CAN CAN'),
        Tedarikci('NU MAĞAZACILIK A.Ş.', '0010000013', 110_000.00, 'YMM CAN CAN'),
        Tedarikci('Ksİ VİNÇ LTD. ŞTİ.', '0010000014', 100_000.00),
        Tedarikci('OMİKRON KİMYA A.Ş.', '0010000015', 90_000.00),
        Tedarikci('PI TRADING GMBH', ITH, 1_200_000.00),
        Tedarikci('RHO METALS LTD', ITH, 130_000.00),
        Tedarikci('SİGMA KIRTASİYE', '0010000018', 30_000.00, takipte=False),
    ]
    return t


@dataclass
class Senaryo:
    donem: Donem
    tedarikciler: list = field(default_factory=varsayilan_tedarikciler)
    r701: tuple = (5_000_000.00, '20', 1_000_000.00)
    yurtici: list = field(default_factory=lambda: [(10_000.00, 1, 100.00), (60_000_000.00, 20, 12_000_000.00)])
    k410: tuple = (40_000_000.00, '20', '4/10', 4_800_000.00)
    k448: tuple = None
    diger: dict = field(default_factory=lambda: {'iade': (100_000.00, 18_000.00), 'kur': (2_000.00, 400.00),
                                                 'diger': (40_000.00, 7_000.00)})
    sorumlu_orani: float = 0.25        # base içindeki sorumlu payı
    devreden_onceki: float = 18_000_000.00
    satis_iade: float = 50_000.00
    i301: tuple = (2_000_000.00, 0.00, 400_000.00)
    i410: tuple = (40_000_000.00, 3_200_000.00)
    i318: tuple = None
    i448: tuple = None
    ihrac_kayitli_bedel: float = 5_000_000.00
    ihrac_kayitli_iade: float = 1_000_000.00
    sonraki_devreden: float = 27_000_000.00
    aylik_bedel: float = 110_000_000.00
    kumulatif: float = 220_000_000.00
    duzeltme: str = ''
    teminat: dict = field(default_factory=lambda: {'301': 360_000.00, '410': 2_880_000.00})
    teminat_tarih: str = '26.03.2026'
    teminat_no: str = '1234567'
    banka: str = 'ÖRNEK KATILIM BANKASI A.Ş. Gaziantep Şubesi'
    liste_farki: float = 0.0           # indirilecek liste ile beyan arasında bırakılacak fark
    liste_donemi: object = None        # None → senaryo dönemi (yyyymm)
    yuklenilen: list = None            # [(firma, vkn, kdv, (yıl, ay))]

    # ---- türetilen tutarlar
    @property
    def ithal(self):
        return sum(t.kdv for t in self.tedarikciler if t.vkn == ITH)

    @property
    def yurtici_ve_sorumlu(self):
        return sum(t.kdv for t in self.tedarikciler if t.vkn != ITH)

    @property
    def sorumlu(self):
        return round(self.yurtici_ve_sorumlu * self.sorumlu_orani, 2)

    @property
    def yurtici_alim(self):
        return round(self.yurtici_ve_sorumlu - self.sorumlu, 2)

    @property
    def base(self):
        return self.yurtici_alim + self.sorumlu + self.ithal

    @property
    def islemler(self):
        s = []
        if self.r701:
            s.append((self.r701[0], self.r701[2]))
        s += [(a, c) for a, _, c in self.yurtici]
        for x in (self.k410, self.k448):
            if x:
                s.append((x[0], x[3]))
        s += list(self.diger.values())
        return s

    @property
    def matrah_toplam(self):
        return round(sum(a for a, _ in self.islemler), 2)

    @property
    def hesaplanan(self):
        return round(sum(b for _, b in self.islemler), 2)

    @property
    def indirim_toplam(self):
        return round(self.devreden_onceki + self.yurtici_alim + self.sorumlu + self.ithal + self.satis_iade, 2)

    @property
    def iade_gereken(self):
        t = 0.0
        if self.i301:
            t += self.i301[2]
        if self.i410:
            t += self.i410[1]
        if self.i318:
            t += self.i318[2]
        if self.i448:
            t += self.i448[1]
        return round(t, 2)

    @property
    def oranlar(self):
        y = self.yurtici_alim
        a, b = round(y * 0.001, 2), round(y * 0.012, 2)
        c = round(y - a - b, 2)
        return [(1, round(a * 100, 2), a), (10, round(b * 10, 2), b), (20, round(c * 5, 2), c)]


AY_PDF = ['Ocak', 'Şubat', 'Mart', 'Nisan', 'Mayıs', 'Haziran', 'Temmuz', 'Ağustos', 'Eylül', 'Ekim', 'Kasım', 'Aralık']


def kdv1_metni(s: Senaryo):
    L = ['KATMA DEĞER VERGİSİ BEYANNAMESİ', 'Vergilendirme Dönemi Tipi Aylık',
         f'Yıl {s.donem.yil}', f'Ay {AY_PDF[s.donem.ay - 1]}', 'MATRAH', 'TEVKİFAT UYGULANMAYAN İŞLEMLER',
         'İşlem Türü Matrah KDV Oranı Vergi']
    if s.r701:
        L.append(f'İhracatı Yapılacak Nihai Ürünlerin Kanunun 11/1-c Maddesi Kapsamında Teslimi {tr(s.r701[0])} {s.r701[1]} {tr(s.r701[2])}')
    for a, b, c in s.yurtici:
        L += ['Yurtiçi Teslim ve Hizmetler', f'{tr(a)} {b} {tr(c)}']
    L.append('KISMİ TEVKİFAT UYGULANAN İŞLEMLER')
    if s.k410:
        L += [f'Yapım İşleri ile Bu İşlerle Birlikte İfa Edilen {tr(s.k410[0])} {s.k410[1]} {s.k410[2]} {tr(s.k410[3])}',
              'Müh.-Mim. Ve Etüt-Proje Hizmetleri [KDVGUT-(I/C-2.1.3.2.1)]']
    if s.k448:
        L += [f'Demir-Çelik Ürünlerinin Teslimi [KDVGUT- {tr(s.k448[0])} {s.k448[1]} {s.k448[2]} {tr(s.k448[3])}', '(I/C-2.1.3.3.8)]']
    L.append('DİĞER İŞLEMLER')
    if 'iade' in s.diger:
        L += ['Alınan Malların İadesi,', f'{tr(s.diger["iade"][0])} {tr(s.diger["iade"][1])}', 'Gerçekleşmeyen İşlemler']
    if 'kur' in s.diger:
        L += ['Kur Farkı / Yuvarlama Farkı', f'{tr(s.diger["kur"][0])} {tr(s.diger["kur"][1])}', 'Nedeniyle Oluşan KDV']
    if 'diger' in s.diger:
        L.append(f'Diğerleri {tr(s.diger["diger"][0])} {tr(s.diger["diger"][1])}')
    if 'amort' in s.diger:
        L += [f'Amortismana Tabi Sabit Kıymet (Taşınmaz, Taşıt Araçları, Demirbaş, {tr(s.diger["amort"][0])} {tr(s.diger["amort"][1])}',
              'Makine ve Teçhizat vb.) Satışları']
    L += [f'Matrah Toplamı {tr(s.matrah_toplam)}', f'Hesaplanan Katma Değer Vergisi {tr(s.hesaplanan)}',
          f'Daha Önce İndirim Konusu Yapılan KDV’nin İlavesi {tr(0)}', f'Toplam Katma Değer Vergisi {tr(s.hesaplanan)}',
          'İNDİRİMLER', f'Önceki Dönemden Devreden {tr(s.devreden_onceki)}', 'İndirilecek KDV',
          f'Yurtiçi Alımlara İlişkin KDV {tr(s.yurtici_alim)}', f'Sorumlu Sıfatıyla Beyan Edilerek Ödenen KDV {tr(s.sorumlu)}']
    if s.ithal:
        L.append(f'İthalde Ödenen KDV {tr(s.ithal)}')
    L += [f'Satıştan İade Edilen, İşlemi Gerçekleşmeyen veya Sonradan İptal Edilen Hizmetler Nedeniyle {tr(s.satis_iade)}',
          'Düzeltilen KDV', f'İndirimler Toplamı {tr(s.indirim_toplam)}',
          'BU DÖNEME AİT TUTİNARDIİNRIİNM OLREARNLARA GÖRE DAĞILIMI', 'KDV Oranı Bedeli KDV Tutarı']
    L += [f'{o} {tr(b)} {tr(v)}' for o, b, v in s.oranlar]
    if s.r701:
        L += ['İHRAÇ KAYDIYLA TESLİMLERE AİT BİLDİRİM',
              f'İhracatı Yapılacak Nihai Ürünlerin Kanunun 11/1-c Maddesi Kapsamında Teslimi {tr(s.r701[0])} {s.r701[1]} {tr(s.r701[2])}',
              f'İhraç Kaydıyla Teslim Bedeli Toplamı {tr(s.ihrac_kayitli_bedel)}', f'Tecil Edilebilir KDV {tr(s.ihrac_kayitli_iade)}',
              f'İhracatın Gerçekleştiği Dönemde İade Edilecek KDV {tr(s.ihrac_kayitli_iade)}']
    L.append('TAM İSTİSNA KAPSAMINA GİREN İŞLEMLER')
    if s.i301:
        L.append(f'301 - Mal İhracatı {tr(s.i301[0])} {tr(s.i301[1])} {tr(s.i301[2])}')
    if s.i318:
        L.append(f'318 - 3996 Sayılı Kanuna Göre Yap İşlet Devret Modeli Çerçevesinde {tr(s.i318[0])} {tr(s.i318[1])} {tr(s.i318[2])}')
    L.append('KISMİ İSTİSNA KAPSAMINA GİREN İŞLEMLER')
    if s.i410:
        L.append(f'410 - Yapım İşleri İle Bu İşlerle Birlikte İfa Edilen {tr(s.i410[0])} {tr(s.i410[1])}')
    if s.i448:
        L.append(f'448 - Demir Çelik Ürünlerinin Teslimi {tr(s.i448[0])} {tr(s.i448[1])}')
    L += ['SONUÇ HESAPLARI', f'Tecil Edilecek Katma Değer Vergisi {tr(0)}', f'Ödenmesi Gereken Katma Değer Vergisi {tr(0)}',
          f'İade Edilmesi Gereken Katma Değer Vergisi {tr(s.iade_gereken)}',
          f'Sonraki Döneme Devreden Katma Değer Vergisi {tr(s.sonraki_devreden)}',
          f'İade Edilebilir KDV {tr(s.iade_gereken)}', 'DİĞER BİLGİLER',
          f'Teslim ve Hizmetlerin Karşılığını Teşkil Eden Bedel (aylık) {tr(s.aylik_bedel)}',
          f'Teslim ve Hizmetlerin Karşılığını Teşkil Eden Bedel (kümülatif) {tr(s.kumulatif)}']
    if s.duzeltme:
        L.append(f'Düzeltme Nedeni : {s.duzeltme}')
    return '\n'.join(L) + '\n'


# ---------------------------------------------------------------- Excel dosyaları (.xls, xlwt)
def _xls(path, satirlar):
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet('Sayfa1')
    for i, r in enumerate(satirlar):
        for j, v in enumerate(r):
            if v is not None:
                ws.write(i, j, v)
    wb.save(str(path))


def _seri(tarih):
    return float((tarih - datetime.date(1899, 12, 30)).days)


def indirilecek_satirlari(s: Senaryo):
    d = s.donem
    donem = s.liste_donemi if s.liste_donemi is not None else float(d.yil * 100 + d.ay)
    bas = ['', 'Sıra No', 'Fatura Tarihi', 'Seri', 'Sıra No', 'Satıcı Adı', 'Satıcı VKN', 'Mal Cinsi', 'Miktar', 'Bedel', 'KDV',
           'Tevkifatsız İndirilen', '2 No.lu Ödenen', 'Toplam İndirilen', 'GÇB No', 'İndirim Dönemi']
    rows = [['İNDİRİLECEK KDV LİSTESİ'], bas]
    n = 1
    toplam_hedef = s.base + s.liste_farki
    ted = list(s.tedarikciler)
    fark = round(toplam_hedef - sum(t.kdv for t in ted), 2)
    for t in ted:
        kdv = t.kdv + (fark if t is ted[-1] else 0)
        parcalar = [round(kdv / t.fatura, 2)] * (t.fatura - 1)
        parcalar.append(round(kdv - sum(parcalar), 2))
        for p in parcalar:
            vkn = float(t.vkn) if t.vkn.isdigit() else t.vkn
            rows.append(['', float(n), _seri(datetime.date(d.yil, d.ay, 5)), 'ABC', f'2026{n:09d}', t.firma, vkn, 'MAL', 1.0,
                         round(p * 5, 2), p, p, 0.0, p, '', donem])
            n += 1
    return rows


def takip_satirlari(s: Senaryo):
    rows = [['', 'SIRA', 'FİRMA ADI', 'KDV TUTARI', 'E-POSTA / AÇIKLAMA', 'SMMM', 'YMM']]
    n = 1
    for t in sorted(s.tedarikciler, key=lambda x: -x.kdv):
        if not t.takipte:
            continue
        rows.append(['', float(n), t.firma, t.kdv, 'muhasebe@ornek.com', 'SMMM X', t.ymm])
        n += 1
    return rows


def yuklenilen_satirlari(s: Senaryo):
    d = s.donem
    yuk = s.yuklenilen
    if yuk is None:
        o = d.onceki()
        yuk = [('GAMA DEMİR SAN. A.Ş.', '0010000003', 300_000.00, (d.yil, d.ay)),
               ('ALFA ÇELİK SAN. VE TİC. A.Ş.', '0010000001', 100_000.00, (o.yil, o.ay))]
    rows = [['', 'SIRA', 'TARİH', 'NO', 'SATICI', 'VKN', '', '', '', '', 'KDV', '', 'KOD', '', '', 'DÖNEM']]
    for i, (f, v, kdv, (yy, mm)) in enumerate(yuk, 1):
        rows.append(['', float(i), _seri(datetime.date(yy, mm, 10)), f'F{i}', f, v, '', '', '', '', kdv, '', '301', '', '', float(yy * 100 + mm)])
    return rows


def dilekce_docx(path, s: Senaryo):
    d = Document()
    d.add_paragraph('GAZİANTEP İHTİSAS VERGİ DAİRESİ MÜDÜRLÜĞÜNE')
    turler = []
    if '301' in s.teminat:
        turler.append(f'301- Mal İhracatından dolayı {tr(s.teminat["301"])} TL')
    if '410' in s.teminat:
        turler.append(f'410- Yapım İşleri İle Bu İşlerle Birlikte İfa Edilen hizmetlerden dolayı {tr(s.teminat["410"])} TL')
    for t in turler:
        d.add_paragraph(t)
    top = sum(s.teminat.values())
    d.add_paragraph(f'Firmamızın {s.donem.egik} dönemine ait KDV iadesi {s.banka} ’nden alınan '
                    f'{s.teminat_tarih} tarih ve {s.teminat_no} nolu teminat mektubu karşılığında toplam {tr(top)} TL talep edilmiştir.')
    d.save(str(path))


def girdi_klasoru(kok, s: Senaryo, ad=None):
    """<kok>/<yyyy>/<NN AY>/ altında bir aylık girdi klasörü üretir ve yolunu döndürür."""
    k = Path(kok) / str(s.donem.yil) / (ad or f'{s.donem.ay:02d} {s.donem.ad}')
    (k / 'TUTANAK ÇALIŞMASI').mkdir(parents=True, exist_ok=True)
    # Gerçek PDF yerine pdfplumber metni yazılır; testlerde pdf_metni bu dosyayı düz metin olarak okur.
    (k / 'KDV 1.pdf').write_text(kdv1_metni(s), encoding='utf-8')
    _xls(k / 'internetvd_kdviadesi_indirilecekkdvListesi_FORMATI.xls', indirilecek_satirlari(s))
    _xls(k / 'TUTANAK ÇALIŞMASI' / f'01 FİRMA VE MUH. BİLGİLERİ {s.donem.ad} {s.donem.yil}.xls', takip_satirlari(s))
    _xls(k / 'yüklenilen tutanak çalışması.xls', yuklenilen_satirlari(s))
    if s.teminat:
        dilekce_docx(k / f'TEMİNAT MEKTUBU KABUL DİLEKÇESİ - {s.donem.ad} {s.donem.yil}.docx', s)
    return k


# ---------------------------------------------------------------- şablon rapor (.docx)
def _tablo(doc, satirlar, sutun):
    """satirlar: [(hücre metinleri, birleşim aralıkları [(a, b)])]."""
    t = doc.add_table(rows=len(satirlar), cols=sutun)
    t.style = 'Table Grid'
    for i, (metin, birlesim) in enumerate(satirlar):
        row = t.rows[i]
        hucreler = []
        for a, b in birlesim or [(j, j) for j in range(sutun)]:
            c = row.cells[a] if a == b else row.cells[a].merge(row.cells[b])
            hucreler.append(c)
        for c, m in zip(hucreler, metin):
            c.text = m
    return t


def _dikey(t, sutun, r0, r1):
    if r1 > r0:
        a = t.cell(r0, sutun)
        a.merge(t.cell(r1, sutun))


def sablon_docx(path, s: Senaryo, rapor_sayisi='2026-10'):
    """s: ŞABLONUN (önceki ayın) senaryosu. Ofis raporunun yapısını taklit eden bir rapor üretir."""
    d, D = s.donem, Document()
    p = D.add_paragraph
    h5 = [(0, 0), (1, 3), (4, 4)]
    kapak = [(['Rapor Sayısı', f'YMM 00000000/{rapor_sayisi}', 'GAZİANTEP'], h5),
             (['Rapor Ekleri', 'Ekler Listesi', '01.04.2026'], h5),
             (['YEMİNLİ MALİ MÜŞAVİRLİK KATMA DEĞER VERGİSİ İADESİ TEMİNAT ÇÖZÜMÜ TASDİK RAPORU'], [(0, 4)]),
             (['Dayanak Sözleşmenin', 'Günü', '04.03.2026'], [(0, 1), (2, 2), (3, 4)]),
             (['Dayanak Sözleşmenin', 'Sayısı', '46'], [(0, 1), (2, 2), (3, 4)]),
             (['Mükellefin', 'Unvanı', 'Örnek Metal San. Tic. A.Ş.'], [(0, 1), (2, 2), (3, 4)]),
             (['İncelemenin Dönemi', d.tire], [(0, 2), (3, 4)]),
             (['SONUÇ', 'Raporun Sonuç Bölümünde Açıklanmıştır.'], [(0, 2), (3, 4)])]
    _tablo(D, kapak, 5)
    p('1-GENEL BİLGİ:')
    p(f'Raporun amacı {d.egik} dönemi KDV beyannamesi kontrolü ve 20.02.2026 tarih ve 7654321 numaralı ESKİ BANKA A.Ş. '
      'Merkez Şubesi’ nin Teminat Mektubuyla alınan KDV İadesinin doğruluğunu araştırmak, tespit etmek ve sağlamaktır.')
    p(f'Firmanın {d.tire} dönemi işlemlerinin KDV iadesi yönünden yeminli mali müşavirliğimce incelenmesi sonucu bu rapor düzenlenmiştir.')
    p('1-1-Ticaret sicil kaydı')
    p('Firma, Gaziantep Ticaret Sicili’ ne 08.07.2009 tarih ve 11111 sicil numarası ile tescillidir. Ticaret Sicil Gazetesi '
      f'15/04/2026 tarih ve YMM 00000000/2026-50 sayılı OCAK-{d.yil} dönemi raporumuza eklendiğinden bu raporumuza eklenmemiştir.')
    p('1-5 Kullanılan Kredilere İlişkin Bilgiler:')
    p(f'Şirketin bir önceki yıl ({d.yil - 1} yılında) kullandığı kredi bilgileri aşağıdaki gibidir.')
    p('1-6 Çalıştırılan İşçi Sayısı')
    p('Firmanın üretim süresi içerisinde çalıştırdığı işçi sayısı aşağıdaki gibidir.')
    aylar = []
    x = d
    for _ in range(12):
        aylar.append(x)
        x = x.onceki()
    aylar.reverse()
    isci = [([aylar[i].tire, str(400 + i), aylar[i + 6].tire, str(450 + i)], None) for i in range(6)]
    _tablo(D, isci, 4)
    p('1-8 En Son Tarihli Bilânço ve Gelir Tablosu Örneği')
    p(f'En son tarihli gelir tablosu 15/04/2026 tarih ve YMM 00000000/2026-50 sayılı OCAK-{d.yil} dönemi raporumuza '
      'eklendiğinden bu raporumuza eklenmemiştir.')
    p('2- USUL İNCELEMELERİ:')
    p(f'Firmanın {d.yil} yılına ait kanuni defterlerinin tasdik durumuna ilişkin bilgiler aşağıdaki gibidir.')
    p('3-1 İLGİLİ DÖNEM KDV BEYANINA İLİŞKİN BİLGİLER:')
    p(f'Şirketin inceleme dönemine ait KDV beyannamesinin fotokopisi ektedir. (Ek:4-{d.yil} {d.ad} KDV Beyannamesi)')
    full5 = [(0, 4)]
    m = [(['MATRAH VE VERGİ BİLDİRİMİ'], full5), (['TEVKİFAT UYGULANMAYAN İŞLEMLER'], full5),
         (['İşlem Türü', 'Matrah', 'KDV Oranı', 'Vergi'], [(0, 0), (1, 1), (2, 2), (3, 4)]),
         (['İhracatı Yapılacak Nihai Ürünlerin Kanunun 11/1-c Maddesi Kapsamında Teslimi', tr(s.r701[0]), '20', tr(s.r701[2])], [(0, 0), (1, 1), (2, 2), (3, 4)]),
         (['Yurtiçi Teslim ve Hizmetler', tr(1.0), '1', tr(0.01)], [(0, 0), (1, 1), (2, 2), (3, 4)]),
         (['Yurtiçi Teslim ve Hizmetler', tr(2.0), '20', tr(0.4)], [(0, 0), (1, 1), (2, 2), (3, 4)]),
         (['KISMİ TEVKİFAT UYGULANAN İŞLEMLER'], full5),
         (['İşlem türü', 'Matrah', 'KDV Oranı', 'Tevkifat', 'Vergi'], None),
         (['Yapım İşleri ile Bu İşlerle Birlikte İfa Edilen Müh.-Mim. Ve Etüt-Proje Hizmetleri [KDVGUT-(I/C-2.1.3.2.1)]', tr(3.0), '20', '4/10', tr(0.36)], None),
         (['DİGER İŞLEMLER'], full5),
         (['İşlem türü', 'Matrah', 'Vergi'], [(0, 2), (3, 3), (4, 4)]),
         (['Alınan Malların İadesi, Gerçekleşmeyen İşlemler', tr(4.0), tr(0.8)], [(0, 2), (3, 3), (4, 4)]),
         (['Diğerleri', tr(5.0), tr(1.0)], [(0, 2), (3, 3), (4, 4)]),
         (['Amortismana Tabi Sabit Kıymet (Taşınmaz, Taşıt Araçları, Demirbaş, Makine ve Teçhizat vb.) Satışları', tr(6.0), tr(0.06)], [(0, 2), (3, 3), (4, 4)]),
         (['Matrah Toplamı', tr(7.0)], [(0, 2), (3, 4)]), (['Hesaplanan KDV Toplamı', tr(8.0)], [(0, 2), (3, 4)]),
         (['Daha önce indirim konusu yapılan KDV’nin İlavesi', tr(0)], [(0, 2), (3, 4)]), (['Toplam KDV', tr(8.0)], [(0, 2), (3, 4)])]
    _tablo(D, m, 5)
    p('3-2 ALIŞ FATURALARINA İLİŞKİN LİSTE')
    p(f'Firmanın inceleme dönemi alışlarına ilişkin liste rapora eklenmiştir. (EK:4-{d.yil} {d.ad} KDV Beyannamesi)')
    _tablo(D, [(['Yurtiçi Alımlara KDV', tr(9.0)], [(0, 1), (2, 2)]), (['Sorumlu Sıfatıyla Beyan Edilen KDV', tr(10.0)], [(0, 1), (2, 2)]),
               (['İthalde Ödenen KDV', tr(11.0)], [(0, 1), (2, 2)]),
               (['BU DÖNEME AİT İNDİRİLECEK KDV TUTARININ ORANLARA GÖRE DAĞILIMI'], [(0, 2)]),
               (['KDV Oranı', 'Alınan Mal ve Hizmete Ait Bedel', 'KDV Tutarı'], None),
               (['%1', tr(12.0), tr(0.12)], None), (['%20', tr(13.0), tr(2.6)], None)], 3)
    p('3-3 YAPILAN İNCELEMELER:')
    p(f'ÖRNEK METAL SAN. TİC. A.Ş. {d.egik} dönemi inceleme döneminde yapılan incelemeler aşağıda açıklanmıştır.')
    p('3-3-1. Mükellefin İş hacmi Durumu')
    p(f'Şirketin {d.egik} dönemi içerisindeki İlgili aylardaki KDV Beyannamesi bilgileri aşağıdaki gibidir.')
    ih = ['TESLİM VE HİZMET TUTARI', 'TOPLAM HES. KDV', 'ÖNCEKİ DÖNEMDEN DEVİR KDV', 'YURTİÇİ ALIMLARINA İLİŞKİN KDV',
          'SORUMLU SIFATIYLA BEYAN EDİLEN KDV', 'İTHALDE ÖDENEN KDV', 'SATIŞTAN İADE EDİLEN KDV', 'TOPLAM İND. KDV',
          'İADE EDİLMESİ GEREKEN KDV', '301 – Mal İhracatı', '410 – Yapım İşleri ile Bu İşlerle Birlikte İfa', 'SONRAKİ DÖNEME DEVREDEN KDV']
    _tablo(D, [([a, tr(s.sonraki_devreden if a.startswith('SONRAKİ') else 14.0 + i)], None) for i, a in enumerate(ih)], 2)
    p('3-4-2 Yapılan Karşıt İncelemelerin Değerlendirilmesi:')
    p('Bakanlıkça belirlenen yeniden değerleme oranı ise; 2023 Yılı İçin VUK. 353 nolu genel tebliğinde karşıt inceleme sınırı '
      '150.000,000 TL bir mükelleften bir aylık dönemde alınan belge toplamı 450.000,00 TL olarak dikkate alınmıştır.')
    p('Yukarıdaki açıklamalara göre, karşıt inceleme yapılacak en az miktar ve karşıt inceleme yapılan miktar aşağıdaki tabloda belirtilmiştir.')
    k = _tablo(D, [([d.tire, 'Firmanın İndirilecek KDV Tutarı', tr(30.0), 'İndirilecek KDV', '%100'], None),
                   (['', '(İndirilecek KDV x %80)', tr(24.0), 'Asgari İnceleme Oranı', '%80'], None),
                   (['', 'Karşıt İnceleme Yapılan Tutar', tr(29.0), 'İnceleme yapılan Oran', '%96,67'], None)], 5)
    _dikey(k, 0, 0, 1)
    k.cell(0, 0).text = d.tire
    p('Firmanın yukarıda yapılan açıklamalar doğrultusunda bu döneme ait indirilecek KDV’sinin yukarıdaki tabloda belirtilen tutardaki kısmı incelenmiştir.')
    p('3-4-3 İndirilecek KDV Karşıt İncelemelerin Safhaları')
    baslik = f'{d.egik} İNDİRİLECEK KDV TUTANAKLARI'
    ek = [(['', 'FİRMANIN ÜNVANI', 'İNCELEME ŞEKLİ', 'RAPOR DURUMU', 'KDV TUTARI', 'İNCELEME ORANI'], [(0, 0), (1, 1), (2, 3), (4, 4), (5, 5), (6, 6)])]
    eski_ekli = [('ESKİ FİRMA BİR A.Ş.', 'BİLGİ İSTEME', '1- EKLİ', 1.0), ('ESKİ FİRMA İKİ A.Ş.', 'K.İ.T.', '2- EKLİ', 2.0),
                 ('ESKİ İTHALAT GMBH', 'İTHALAT', '3- İTHALAT', 3.0)]
    for f, sk, du, v in eski_ekli:
        ek.append((['', f, sk, du, tr(v), '% 1,00'], [(0, 0), (1, 1), (2, 3), (4, 4), (5, 5), (6, 6)]))
    ek.append((['', 'İNCELEME ORANI', 'TOPLAM', tr(6.0), '% 3,00'], [(0, 0), (1, 2), (3, 4), (5, 5), (6, 6)]))
    m0 = len(ek)
    for f, sk, v in [('ESKİ MUHAFAZA BİR', 'K.İ.T.', 4.0), ('ESKİ MUHAFAZA İKİ', 'BİLGİ İSTEME', 5.0)]:
        ek.append((['', f, sk, 'Tarafımızdan Muhafaza Edilmektedir.', tr(v), '% 2,00'], [(0, 0), (1, 1), (2, 3), (4, 4), (5, 5), (6, 6)]))
    m1 = len(ek) - 1
    ek.append((['', 'TOPLAM', tr(9.0), '% 2,00'], [(0, 0), (1, 4), (5, 5), (6, 6)]))
    ek.append((['', 'GENEL TOPLAM İNCELEME ORANI', tr(15.0), '% 5,00'], [(0, 0), (1, 4), (5, 5), (6, 6)]))
    t12 = _tablo(D, ek, 7)
    _dikey(t12, 6, m0, m1)
    t12.cell(m0, 6).text = '% 2,00'
    _dikey(t12, 0, 0, len(ek) - 1)
    t12.cell(0, 0).text = baslik
    p('3-4-4 Yüklenilen KDV Karşıt İncelemelerin Safhaları')
    o = d.onceki()
    yk = [(['', 'FİRMANIN ÜNVANI', 'İNCELEME ŞEKLİ', 'İNCELEME DÖNEMİ', 'RAPOR DURUMU', 'ASIL TUTANAKLARIN DURUMU', 'KDV TUTARI', 'İNCELEME ORANI'], None),
          (['', 'GAMA DEMİR SANAYİ A.Ş.', 'BİLGİ İSTEME', o.egik, '1- EKLİ', '01.01.2026 – YMM 00000000/2026-1 Sayılı Raporda Mevcut', tr(1.0), '% 50,00'], None),
          (['', 'ESKİ YÜKLENİLEN A.Ş.', 'BİLGİ İSTEME', o.egik, '2- EKLİ', '01.01.2026 – YMM 00000000/2026-1 Sayılı Raporda Mevcut', tr(1.0), '% 50,00'], None),
          (['', 'TOPLAM İNCELEME ORANI', tr(2.0), '% 100,00'], [(0, 0), (1, 5), (6, 6), (7, 7)])]
    t13 = _tablo(D, yk, 8)
    _dikey(t13, 0, 0, len(yk) - 1)
    t13.cell(0, 0).text = 'YÜKLENİLEN KDV TUTANAKLARI'
    p('3.5. İhraç Kaydıyla Kesilen Faturaların Gerçek Olup Olmadığının Tespiti:')
    p(f'Firmanın {d.egik} dönemi ihracatı yapılacak nihai ürünlerin kanunun 11/1-c maddesi kapsamında toplam {tr(1.5)} TL teslimde '
      f'bulunulmuştur. Yapılan teslimden dolayı toplam {tr(0.3)} TL KDV iadesi hesaplanmıştır ve bu tutar mahsup dosyası ile talep edilmiştir.')
    p('3-6. 301- Mal İhracatı Kapsamında Kesilen Satış faturalarının Gerçek Olup Olmadığının Tespiti:')
    p(f'Firmanın {d.egik} dönemi mal ve/veya hizmet ihracatından dolayı toplam {tr(2.5)}- TL teslimlerde bulunmuştur. '
      f'Yapılan ihracattan dolayı toplam {tr(0.5)}- TL KDV iadesi hesaplanmıştır.')
    p('3.7. Tevkifatlı Satış Faturalarının Gerçek Olup Olmadığının Tespiti:')
    p(f'Firmanın {d.egik} Dönemindeki tevkifatlı satış bilgileri aşağıdaki gibidir.')
    _tablo(D, [(['', 'KISMİ TEVKİFAT KAPSAMINA GİREN İŞLEMLER'], [(0, 0), (1, 5)]),
               (['Dönemi', 'Teslim Bedeli', 'Oran', 'Hesaplanan KDV', 'Tevkifat Oranı', 'İadeye Konu KDV'], None),
               ([d.tire, tr(3.0), '%20', tr(0.6), '4/10', tr(0.24)], None)], 6)
    p(f'Firmanın {d.egik} dönemine ait YAPIM İŞLERİ (410) bölümündeki satışlarına istinaden iade edilebilir tutar {tr(0.24)} TL olarak hesaplanmıştır.')
    p('3-8-1 KDV Beyannamesi ile İlgili Düzeltme Bilgisi;')
    p(f' Firmanın {d.egik} dönemine KDV beyannamesine düzeltme verilmemiştir.')
    f4 = [(0, 4)]
    r3 = [(0, 2), (3, 3), (4, 4)]
    r2 = [(0, 3), (4, 4)]
    r4 = [(0, 0), (1, 1), (2, 2), (3, 4)]
    dk = [([f'MATRAH VE VERGİ BİLDİRİMİ-{d.egik}'], f4), (['', 'TEVKİFAT UYGULANMAYAN İŞLEMLER'], [(0, 0), (1, 4)]),
          (['İşlem Türü', 'Matrah', 'KDV Oranı', 'Vergi'], r4),
          (['İhracatı Yapılacak Nihai Ürünlerin Kanunun 11/1-c Maddesi Kapsamında Teslimi', tr(1.0), '20', tr(0.2)], r4),
          (['Yurtiçi Teslim ve Hizmetler', tr(1.0), '20', tr(0.2)], r4),
          (['KISMİ TEVKİFAT UYGULANAN İŞLEMLER'], f4), (['İşlem türü', 'Matrah', 'KDV Oranı', 'Tevkifat', 'Vergi'], None),
          (['Yapım İşleri ile Bu İşlerle Birlikte İfa Edilen Müh.-Mim. Ve Etüt-Proje Hizmetleri [KDVGUT-(I/C-2.1.3.2.1)]', tr(1.0), '20', '4/10', tr(0.08)], None),
          (['Demir-Çelik Ürünlerinin Teslimi [KDVGUT-(I/C-2.1.3.3.8)]', tr(1.0), '20', '5/10', tr(0.1)], None),
          (['DİĞER İŞLEMLER'], f4), (['İşlem türü', 'Matrah', 'Vergi'], r3),
          (['Kur Farkı / Yuvarlama Farkı Nedeniyle Oluşan KDV', tr(1.0), tr(0.2)], r3),
          (['Matrah Toplamı', tr(1.0)], r2), (['Hesaplanan KDV Toplamı', tr(1.0)], r2), (['İlave Edilecek KDV', tr(0)], r2),
          (['Toplam KDV', tr(1.0)], r2), (['DİĞER İNDİRİMLER'], f4), (['Satıştan İade Edilen KDV', tr(1.0)], r2),
          (['Yurtiçi Alımlarına İlişkin KDV', tr(1.0)], r2), (['Sorumlu Sıfatıyla Beyan Edilen KDV', tr(1.0)], r2),
          (['İthalde Ödenen KDV', tr(1.0)], r2), (['ÖNCEKİ DÖNEMDEN DEVREDEN İNDİRİLECEK KDV'], f4),
          (['Önceki Dönemden Devreden KDV', tr(1.0)], r2), (['İndirimler Toplamı', tr(1.0)], r2),
          (['BU DÖNEME AİT İNDİRİLECEK KDV TUTARININ ORANLARA GÖRE DAĞILIMI'], f4),
          (['KDV Oranı', 'Alınan Mal ve Hizmete Ait Bedel', 'KDV Tutarı'], r3), (['%1', tr(1.0), tr(0.01)], r3), (['%20', tr(1.0), tr(0.2)], r3),
          (['İHRAÇ KAYDIYLA TESLİMLERE AİT BİLDİRİM'], f4), (['İşlem Türü', 'Teslim Bedeli', 'Oran', 'Hesaplanan KDV'], r4),
          (['İhracatı Yapılacak Nihai Ürünlerin Kanunun 11/1-c Maddesi Kapsamında Teslimi', tr(1.0), '20', tr(0.2)], r4),
          (['İhraç Kaydıyla Teslim Bedeli Toplamı', tr(1.0)], r2), (['Tecil Edilebilir KDV', tr(1.0)], r2),
          (['İhracatın Gerçekleştiği Dönemde İade Edilecek KDV', tr(1.0)], r2),
          (['TAM İSTİSNA KAPSAMINA GİREN İŞLEMLER'], f4), (['İstisna Türü', 'Teslim ve Hizmetler Tutarı', 'İadeye Konu KDV'], r3),
          (['301-Mal İhracatı', tr(1.0), tr(0.2)], r3),
          (['KISMİ TEVKİFAT KAPSAMINA GEREN İŞLEMLER'], f4), (['İstisna Türü', 'Teslim ve Hizmet Tutarı', 'İadeye Konu Olan KDV'], r3),
          (['410-Yapım İşleri ile Bu İşlerle Birlikte İfa', tr(1.0), tr(0.2)], r3),
          (['448-Demir Çelik Ürünlerinin Teslimi', tr(1.0), tr(0.0)], r3),
          (['Tecil Edilecek Katma Değer Vergisi', tr(0)], r2), (['Ödenmesi Gereken Katma Değer Vergisi', tr(0)], r2),
          (['İade Edilmesi Gereken Katma Değer Vergisi', tr(1.0)], r2),
          (['Sonraki Döneme Devreden Katma Değer Vergisi', tr(s.sonraki_devreden)], r2),
          (['Teslim ve Hizmetlerin Karşılığını Teşkil Eden Bedel (Aylık)', tr(1.0)], r2),
          (['Teslim ve Hizmetlerin Karşılığını Teşkil Eden Bedel (Kümülâtif)', tr(1.0)], r2)]
    _tablo(D, dk, 5)
    p('4. SONUÇ:')
    p(f'4-3. {d.egik} KDV Beyanına göre 701- İhracatı Yapılacak Nihai Ürünlerin Kanunun 11/1-c Maddesi Kapsamında Tesliminden dolayı '
      f'{tr(0.3)} TL KDV iadesi hesaplanmıştır ve bu tutar mahsup dosyası ile talep edilmiştir.')
    p(f'4-4. {d.egik} KDV Beyanına göre 301-Mal İhracatından dolayı {tr(0.5)} TL, 410- Yapım İşleri İle Bu İşlerle Birlikte İfa '
      f'Edilen Mühendislik-Mimarlık ve Etüt-Proje hizmetlerinde Tevkifata Tabi Tutulan dolayı {tr(0.24)} TL KDV iadesi hesaplanmıştır.')
    p(f'4-5. {d.egik} KDV Beyanına göre 301-Mal İhracatından dolayı {tr(0.4)} TL, 410- Yapım İşleri İle Bu İşlerle Birlikte İfa '
      f'Edilen Mühendislik-Mimarlık ve Etüt-Proje hizmetlerinde Tevkifata Tabi Tutulan {tr(0.2)} TL KDV’nin iade edilmesi gerektiği ve '
      'mükellefe 20.02.2026 tarih ve 7654321 numaralı ESKİ BANKA A.Ş. Merkez Şubesi’ nin Teminat Mektubuyla toplam '
      f'{tr(0.6)} TL iade edildiğinden söz konusu Teminat Mektubunun serbest bırakılması,')
    p(f'4-6. {d.egik} döneminde teminat mektubu ile alınamayan ve iade edilmesi gereken KDV’nin {tr(0.14)}- TL olduğu,')
    p(f'4-8. {d.egik} KDV Beyanına göre sonraki döneme devir olan KDV tutarının {tr(s.sonraki_devreden)}--TL olarak dikkate alınması,')
    D.save(str(path))
    return Path(path)


def senaryo_kopya(s, **degisiklik):
    t = copy.deepcopy(s)
    for k, v in degisiklik.items():
        setattr(t, k, v)
    return t


# ---------------------------------------------------------------- gerçek PDF (reportlab) — .exe duman testi için
YAZI_TIPLERI = [r'C:\Windows\Fonts\arial.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
                '/System/Library/Fonts/Supplemental/Arial.ttf']


def yazi_tipi():
    return next((p for p in YAZI_TIPLERI if Path(p).is_file()), None)


def kdv1_pdf(path, s: Senaryo):
    """KDV 1 metnini satır satır gerçek bir PDF'e yazar (Türkçe karakterli TTF yazı tipiyle)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    yt = yazi_tipi()
    if not yt:
        raise RuntimeError('Türkçe karakterli yazı tipi bulunamadı')
    pdfmetrics.registerFont(TTFont('Sentetik', yt))
    c = canvas.Canvas(str(path), pagesize=A4)
    y = A4[1] - 30
    for satir in kdv1_metni(s).splitlines():
        if y < 30:
            c.showPage()
            y = A4[1] - 30
        c.setFont('Sentetik', 8)
        c.drawString(20, y, satir)
        y -= 11
    c.save()
    return Path(path)
