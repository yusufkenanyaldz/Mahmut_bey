import pytest
from docx import Document

from teminat.rapor import docx_araclari as dx


def test_paragraf_degistir_bicimi_korur():
    d = Document()
    p = d.add_paragraph()
    p.add_run('Firmanın ŞUBAT/2026 dönemi toplam ')
    kalin = p.add_run('1.234,56')
    kalin.bold = True
    p.add_run(' TL teslimde bulunmuştur.')
    assert dx.paragraf_degistir(p, r'[\d.]+,\d\d', '9.876,54')
    assert p.text == 'Firmanın ŞUBAT/2026 dönemi toplam 9.876,54 TL teslimde bulunmuştur.'
    assert len(p.runs) == 3 and p.runs[1].bold and p.runs[1].text == '9.876,54'


def test_paragraf_degistir_runlara_yayilan_eslesme():
    d = Document()
    p = d.add_paragraph()
    p.add_run('ŞUBAT')
    p.add_run('/2026 dönemi')
    assert dx.metin_degistir(p, 'ŞUBAT/2026', 'MART/2026')
    assert p.text == 'MART/2026 dönemi'
    assert not dx.metin_degistir(p, 'NİSAN/2026', 'x')


def _tablo(satirlar):
    d = Document()
    t = d.add_table(rows=len(satirlar), cols=2)
    for r, (a, b) in zip(t.rows, satirlar):
        r.cells[0].text, r.cells[1].text = a, b
    return t


def test_grubu_esitle_ekler_siler_siralar():
    t = _tablo([('BAŞLIK', ''), ('410 – Yapım', '1,00'), ('448 – Demir', '2,00'), ('SONRAKİ', '3,00')])
    mevcut = [t.rows[1], t.rows[2]]
    istenen = [('301', 5.0), ('410', 6.0)]
    sonuc, silinen = dx.grubu_esitle(mevcut, istenen, lambda x: x[0], lambda r: r.cells[0].text[:3])
    assert [(k[0], yeni) for _, k, yeni in sonuc] == [('301', True), ('410', False)]
    assert silinen == ['448 – Demir']
    for r, (kod, v), yeni in sonuc:
        if yeni:
            dx.set_cell(r.cells[0], f'{kod} – Yeni')
        dx.set_row_vals(r, [f'{v:.2f}'])
    assert [dx.row_text(r) for r in t.rows] == ['BAŞLIK | ', '301 – Yeni | 5.00', '410 – Yapım | 6.00', 'SONRAKİ | 3,00']


def test_grubu_esitle_bos_grup_sablonsuz():
    t = _tablo([('BAŞLIK', ''), ('SON', '')])
    with pytest.raises(dx.SablonYok):
        dx.grubu_esitle([], [('301', 1.0)], lambda x: x[0], lambda r: None, sonra=t.rows[0])
    sonuc, _ = dx.grubu_esitle([], [('301', 1.0)], lambda x: x[0], lambda r: None, sonra=t.rows[0], sablon=t.rows[1])
    assert len(t.rows) == 3 and sonuc[0][2]


def test_set_row_vals_etiketin_ustune_yazmaz():
    t = _tablo([('Etiket', '1,00')])
    with pytest.raises(dx.SablonYok):
        dx.set_row_vals(t.rows[0], ['1', '2'])


def test_ara_turkce_ve_buyuk_kucuk_harf_duyarsiz():
    assert dx.ara(r'^DİĞER İŞLEMLER', 'DİGER İŞLEMLER')
    assert dx.ara(r'^Sonraki Döneme', 'SONRAKİ DÖNEME DEVREDEN')
