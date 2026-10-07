"""İçerikten belge tanıma (PROJE_TALIMATI.md §13): dosya adı ve klasör düzeni ne olursa olsun doğru belgeler seçilmeli."""
import shutil
import unicodedata

import pytest

import sentetik as S
from teminat.islem import hazirla, yol_temizle
from teminat.klasorler import ay_coz
from teminat.ortak import Donem, excel_satirlari
from teminat.tanima import (BelirsizSecim, TanimaHatasi, TurDuzeltmeleri, ad_anahtari, metindeki_donemler, secim_yap,
                            siniflandir, tani)

SUBAT = Donem(2026, 2)


@pytest.mark.parametrize('ad,beklenen', [
    ('02 ŞUBAT', (2, None)), ('02 ŞUBAT (GİRDİ)', (2, None)), ('ABC 12- ARALIK', (12, None)), ('Şubat 2026', (2, 2026)),
    ('2026-02', (2, 2026)), ('2026 02', (2, 2026)), ('02 ŞUBAT (20.03.2026 VERİLDİ)', (2, 2026)), ('3- GİRDİ', (3, None)),
    (unicodedata.normalize('NFD', '02 ŞUBAT'), (2, None)), ('SİSTEM', None), ('OCAK-ŞUBAT', None), ('MARTI YAPI', None),
])
def test_klasor_adindan_ay_ipucu(ad, beklenen):
    assert ay_coz(ad) == beklenen


@pytest.mark.parametrize('metin,tur', [
    ('KATMA DEĞER VERGİSİ BEYANNAMESİ\nGerçek Usulde Katma Değer Vergisi Mükellefleri İçin', 'KDV1'),
    ('KATMA DEĞER VERGİSİ BEYANNAMESİ Vergi Sorumluları İçin', 'KDV2'),
    ('Sıra No Alış Faturasının Tarihi Seri Sıra Satıcının Adı Soyadı Toplam İndirilen KDV Tutarı', 'LISTE_INDIRILECEK'),
    ('SIRA FİRMA ADI KDV TUTARI E-POSTA / AÇIKLAMA SMMM YMM', 'TAKIP'),
    ('Firmamızın ŞUBAT/2026 dönemi … teminat mektubu karşılığında toplam 3.240.000,00 TL talep edilmiştir.', 'TEMINAT_DILEKCE'),
    ('1-GENEL BİLGİ:\nRaporun amacı ŞUBAT/2026 dönemi', 'RAPOR'),
    ('', 'TARANMIS'), ('e-Arşiv Fatura No ABC2026', 'BILINMEYEN'),
])
def test_siniflandirma_kurallari(metin, tur):
    assert siniflandir(metin) == tur


def test_karisik_klasor_icerikten_taninir(tmp_path):
    subat, v = S.karisik_firma(tmp_path)
    belgeler = tani(subat)
    tur = {b.goreli: b.tur for b in belgeler}
    assert tur['beyan_subat.pdf'] == 'KDV1' and tur['kdv2.pdf'] == 'KDV2'
    assert tur['liste.xls'] == 'LISTE_INDIRILECEK' and tur['SİSTEM/liste.xls'] == 'LISTE_INDIRILECEK'
    assert tur['muh bilgi.xls'] == 'TAKIP' and tur['çalışma/yük.xls'] == 'LISTE_YUKLENILEN'
    assert tur['dilekce.docx'] == 'TEMINAT_DILEKCE' and tur['eski/dilekce ocak.docx'] == 'TEMINAT_DILEKCE'
    assert tur['tutanak/alfa.docx'] == 'KIT' and tur['tutanak/yazı.docx'] == 'YMM_YAZISI' and tur['imalat.xlsx'] == 'IMALAT'
    assert tur['301/tarama.jpg'] == 'TARANMIS'
    assert '._beyan_subat.pdf' not in tur                         # macOS çöpü atlanır
    s = secim_yap(subat, belgeler)
    assert s.donem == SUBAT
    assert {r: p.relative_to(subat).as_posix() for r, p in s.yollar.items()} == {
        'KDV1': 'beyan_subat.pdf', 'LISTE_INDIRILECEK': 'liste.xls', 'TAKIP': 'muh bilgi.xls',
        'TEMINAT_DILEKCE': 'dilekce.docx', 'LISTE_YUKLENILEN': 'çalışma/yük.xls'}
    assert not s.eksik_zorunlu()


def test_sablon_icerikten_ve_vkn_ile_bulunur(tmp_path):
    subat, v = S.karisik_firma(tmp_path)
    hz = hazirla(subat)
    assert hz.sablon == v['sablon'].resolve()                      # başka firmanın aynı dönem raporu alınmadı
    assert hz.sablon_kaynagi == 'içerikten'
    assert any('başka mükellefe' in n[2] for n in hz.notlar)
    tablo = hz.tablo()
    assert '✔ Şablon (önceki ayın raporu, içerikten)' in tablo and '✘ EKSİK' not in tablo


def test_karisik_klasorden_taslak(tmp_path, ayar):
    from teminat.taslak import taslak_uret
    subat, v = S.karisik_firma(tmp_path)
    tc = taslak_uret(ayar, None, subat)
    assert tc.docx.exists() and tc.docx.parent == subat.resolve() / 'CLAUDE TASLAK'
    assert tc.sonuc.kontroller.say('HATA') == 0, [(k.konu, k.aciklama) for k in tc.sonuc.kontroller if k.durum == 'HATA']
    from openpyxl import load_workbook
    wb = load_workbook(tc.xlsx)
    assert 'Belgeler' in wb.sheetnames


def test_bu_ayin_raporu_varsa_karsilastirilir(tmp_path, ayar):
    from teminat.taslak import taslak_uret
    subat, v = S.karisik_firma(tmp_path)
    ilk = taslak_uret(ayar, None, subat, tmp_path / 'ilk')
    shutil.copy(ilk.docx, subat / 'ofisin raporu.docx')             # ofisin bitmiş raporu gibi
    tc = taslak_uret(ayar, None, subat, tmp_path / 'ikinci')
    assert tc.farklar and tc.karsilastirma.tutar_farki == 0


def test_klasor_adi_baska_ay_derse_uyari(tmp_path):
    g = S.girdi_klasoru_yol(tmp_path / '05 MAYIS', S.Senaryo(Donem(2026, 3)))
    s = secim_yap(g, tani(g))
    assert s.donem == Donem(2026, 3)
    assert any(n[1] == 'Belge dönemi' and 'MART-2026' in n[2] for n in s.notlar)


def test_ayni_donemin_farkli_iki_beyannamesi_sorulur(tmp_path):
    s1 = S.Senaryo(SUBAT)
    g = S.girdi_klasoru_yol(tmp_path / 'x', s1)
    (g / 'duzeltme.pdf').write_text(S.kdv1_metni(S.senaryo_kopya(s1, duzeltme='Hatalı beyan')), encoding='utf-8')
    with pytest.raises(BelirsizSecim) as e:
        secim_yap(g, tani(g))
    assert e.value.rol == 'KDV1' and len(e.value.adaylar) == 2
    s = secim_yap(g, tani(g), zorla={'KDV1': g / 'duzeltme.pdf'})
    assert s.k['duzeltme'] == 'Hatalı beyan'


def test_iki_donemin_beyannamesi_varsa_en_sonuncusu(tmp_path):
    g = S.girdi_klasoru_yol(tmp_path / 'x', S.Senaryo(Donem(2026, 3)))
    (g / 'onceki.pdf').write_text(S.kdv1_metni(S.Senaryo(SUBAT)), encoding='utf-8')
    s = secim_yap(g, tani(g))
    assert s.donem == Donem(2026, 3) and any('birden fazla dönemin' in n[2] for n in s.notlar)
    assert secim_yap(g, tani(g), istenen=SUBAT).donem == SUBAT


def test_beyanname_yoksa_anlasilir_hata(tmp_path):
    g = S.girdi_klasoru_yol(tmp_path / 'x', S.Senaryo(SUBAT))
    (g / 'KDV 1.pdf').unlink()
    with pytest.raises(TanimaHatasi, match='1 No.lu KDV beyannamesi bulunamadı'):
        secim_yap(g, tani(g))


def test_elle_tur_duzeltmesi_sonraki_ay_da_gecerli(tmp_path):
    dosya = tmp_path / 'belge_turleri.yaml'
    d = TurDuzeltmeleri(dosya)
    d.ayarla(tmp_path / '01 FİRMA VE MUH. BİLGİLERİ ŞUBAT 2026.xls', 'YOK_SAY')
    d2 = TurDuzeltmeleri(dosya)
    assert d2.tur(tmp_path / '01 FİRMA VE MUH. BİLGİLERİ MART 2026.xls') == 'YOK_SAY'
    assert ad_anahtari('liste ŞUBAT 2026.xls') == ad_anahtari('liste Mart 2027.xls')
    with pytest.raises(ValueError):
        d.ayarla(tmp_path / 'a.xls', 'UYDURMA')


def test_elle_duzeltme_secimi_degistirir(tmp_path):
    g = S.girdi_klasoru_yol(tmp_path / 'x', S.Senaryo(SUBAT))
    duz = TurDuzeltmeleri(tmp_path / 'd.yaml')
    duz.ayarla(g / 'yüklenilen tutanak çalışması.xls', 'YOK_SAY')
    b = {x.goreli: x for x in tani(g, duz)}
    assert b['yüklenilen tutanak çalışması.xls'].tur == 'YOK_SAY' and b['yüklenilen tutanak çalışması.xls'].elle


def test_xls_uzantili_xlsx_okunur(tmp_path):
    p = S._xlsx_yaz(tmp_path / 'a.xls', [['başlık'], ['', 1.0, 'x']])
    assert [r for _, r in excel_satirlari(p)][1][:3] == [None, 1, 'x']
    with pytest.raises(ValueError, match='Excel dosyası değil'):
        list(excel_satirlari(S._docx_yaz(tmp_path / 'b.xls', ['metin'])))


def test_metindeki_donemler():
    c = metindeki_donemler('Firmamızın ŞUBAT/2026 dönemine ait; 2026/Şubat; 02/2026 ve Mart 2025')
    assert c[SUBAT] == 3 and c[Donem(2025, 3)] == 1


def test_yol_temizle(tmp_path):
    assert yol_temizle(f'  "{tmp_path}"  ') == tmp_path.resolve()
    assert yol_temizle('   ') is None


def test_cok_dosyali_klasor_reddedilir(tmp_path, monkeypatch):
    from teminat import tanima
    monkeypatch.setattr(tanima, 'COK_DOSYA', 5)
    for i in range(7):
        (tmp_path / f'{i}.txt').write_text('x')
    with pytest.raises(TanimaHatasi, match='ayın klasörü gibi görünmüyor'):
        tanima.dosyalari_listele(tmp_path, tanima.COK_DOSYA)


def test_gercek_pdf_ile_tanima(tmp_path):
    pytest.importorskip('reportlab')
    if not S.yazi_tipi():
        pytest.skip('Türkçe karakterli yazı tipi yok')
    subat, _ = S.karisik_firma(tmp_path, gercek_pdf=True)
    s = secim_yap(subat, tani(subat))
    assert s.donem == SUBAT and s.yollar['KDV1'].name == 'beyan_subat.pdf'
