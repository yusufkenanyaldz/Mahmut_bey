"""Uçtan uca: sentetik ŞUBAT-2026 şablonu + MART-2026 girdileri → taslak + kontrol listesi."""
import hashlib

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

import sentetik as S
from teminat.okuyucular.girdiler import girdileri_oku
from teminat.ortak import Donem, tr
from teminat.rapor.docx_araclari import row_text, uniq_cells
from teminat.rapor.olustur import SablonHatasi, taslak_olustur
from teminat.rapor.tablolar import tablolari_bul
from teminat.taslak import taslak_uret


def _uret(ay, ayar, tmp_path, **degisiklik):
    yeni = S.senaryo_kopya(ay['yeni'], **degisiklik)
    girdi = S.girdi_klasoru(tmp_path / 'GIRDI', yeni)
    out = tmp_path / 'taslak.docx'
    sonuc = taslak_olustur(ay['sablon'], girdileri_oku(girdi), ayar, out, kati=True)
    return Document(str(out)), sonuc, yeni


def _durumlar(sonuc, konu=None):
    return [(k.durum, k.konu, k.aciklama) for k in sonuc.kontroller if konu is None or k.konu == konu]


def _metin(d):
    return '\n'.join(p.text for p in d.paragraphs)


def _sari(p):
    return any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in p.runs)


def test_temel_taslak(ay, ayar, tmp_path):
    sablon_ozet = hashlib.sha256(ay['sablon'].read_bytes()).hexdigest()
    d, sonuc, s = _uret(ay, ayar, tmp_path)
    assert hashlib.sha256(ay['sablon'].read_bytes()).hexdigest() == sablon_ozet   # şablona yazılmaz
    T, sorunlar = tablolari_bul(d)
    assert not sorunlar
    metin = _metin(d)
    assert 'MART/2026' in metin and 'ŞUBAT/2026' not in metin and 'ŞUBAT-2026' not in metin
    assert '26.03.2026 tarih ve 1234567 numaralı' in metin
    assert '15/04/2026 tarih ve YMM 00000000/2026-50 sayılı OCAK-2026 dönemi raporumuza eklendiğinden' in metin
    # kapak
    kap = [row_text(r) for r in T['kapak'].rows]
    assert 'Rapor Sayısı | YMM 00000000/[DOLDURULACAK] | GAZİANTEP' in kap and 'İncelemenin Dönemi | MART-2026' in kap
    # işçi tablosu kaydı
    assert row_text(T['isci'].rows[-1]).endswith('MART-2026 | [DOLDURULACAK]')
    # karşıt inceleme özeti
    k0 = [c.text for c in uniq_cells(T['karsit'].rows[0])]
    assert k0[0] == 'MART-2026' and k0[2] == tr(s.base)
    # safha tablosu
    satir = [row_text(r) for r in T['safha'].rows]
    assert satir[1].endswith(f'ALFA ÇELİK SAN. VE TİC. A.Ş. | BİLGİ İSTEME | 1- EKLİ | {tr(5_000_000)} | % 25,25')
    assert any('12- İTHALAT' in x for x in satir) and 'MART/2026 İNDİRİLECEK KDV TUTANAKLARI' in satir[0]
    assert sum('Muhafaza' in x for x in satir) == 5
    # yüklenilen: önceki dönem referansı ayardan, bu dönemin referansı [DOLDURULACAK]
    y = [row_text(r) for r in T['yuklenilen'].rows]
    assert any('ŞUBAT/2026 | 2- EKLİ | 20.05.2026 – YMM 00000000/2026-60 Sayılı Raporda Mevcut' in x for x in y)
    assert any('MART/2026 | 1- EKLİ | [DOLDURULACAK]' in x for x in y)
    # 4-x paragrafları
    assert f'KDV’nin {tr(s.iade_gereken - 3_240_000)}- TL' in metin
    assert f'tutarının {tr(s.sonraki_devreden)}--TL' in metin
    assert 'toplam 3.240.000,00 TL iade' in metin
    # kontrol listesi: hata yok, beklenen ELLE kalemleri var
    assert sonuc.kontroller.say('HATA') == 0, _durumlar(sonuc)
    konular = {k.konu for k in sonuc.kontroller}
    assert {'Kapak', 'İşçi sayısı', 'Devreden KDV', 'Safha tablosu', 'Eski metin', 'Liste dönemi'} <= konular


def test_448_ve_amortisman_eklenir_sablondaki_fazlalik_silinir(ay, ayar, tmp_path):
    d, sonuc, s = _uret(ay, ayar, tmp_path, k448=(300_000.0, '20', '5/10', 30_000.0),
                        diger={'kur': (2_000.0, 400.0), 'amort': (600_000.0, 6_000.0)}, i448=(300_000.0, 0.0))
    T, _ = tablolari_bul(d)
    m = [row_text(r) for r in T['matrah'].rows]
    assert any(x.startswith('Demir-Çelik Ürünlerinin Teslimi') and tr(30_000) in x for x in m)
    dig = [x.split(' | ')[0] for x in m[m.index('DİGER İŞLEMLER') + 2:m.index(next(x for x in m if x.startswith('Matrah Toplamı')))]]
    assert dig[0].startswith('Kur Farkı') and dig[-1].startswith('Amortismana Tabi')          # amortisman en sonda
    assert not any(x.startswith('Alınan Malların İadesi') or x.startswith('Diğerleri') for x in m)
    dk = [row_text(r) for r in T['dokum'].rows]
    assert sum(x.startswith('Demir-Çelik') for x in dk) == 1                                  # eski 448 satırı kopya kalmaz
    assert any(x.startswith('448-Demir Çelik') and x.endswith('0,00') for x in dk)
    # iş hacminde iadesi 0 olan 448 yazılmaz
    assert not any(row_text(r).startswith('448') for r in T['is_hacmi'].rows)
    assert sonuc.kontroller.say('HATA') == 0, _durumlar(sonuc)


def test_yeni_iade_turu_is_hacmine_sari_eklenir(ay, ayar, tmp_path):
    d, sonuc, s = _uret(ay, ayar, tmp_path, i448=(300_000.0, 10_000.0))
    T, _ = tablolari_bul(d)
    r = next(r for r in T['is_hacmi'].rows if row_text(r).startswith('448'))
    assert row_text(r) == f'448 – Demir Çelik Ürünlerinin Teslimi | {tr(10_000)}'
    assert _sari(uniq_cells(r)[0].paragraphs[0])
    assert any(k.konu == 'Yeni satır' and '448' in k.aciklama for k in sonuc.kontroller)


def test_teminat_dilekcesi_yoksa_hata_ve_sari(ay, ayar, tmp_path):
    d, sonuc, _ = _uret(ay, ayar, tmp_path, teminat={})
    hatalar = _durumlar(sonuc)
    assert ('HATA', 'Teminat mektubu') in [(a, b) for a, b, _ in hatalar]
    assert ('HATA', 'Teminat dilekçesi') in [(a, b) for a, b, _ in hatalar]
    for p in d.paragraphs:
        if 'Teminat Mektub' in p.text:
            assert _sari(p) and '20.02.2026 tarih ve 7654321' in p.text     # eski tarih kaldı ama işaretli


def test_liste_donemi_ve_toplam_farki(ay, ayar, tmp_path):
    _, sonuc, _ = _uret(ay, ayar, tmp_path, liste_donemi=202602.0, liste_farki=1234.56)
    d = {b: (a, c) for a, b, c in _durumlar(sonuc)}
    assert d['Liste dönemi'][0] == 'HATA' and 'ŞUBAT/2026: ' in d['Liste dönemi'][1]
    assert d['İndirilecek liste ↔ beyan'][0] == 'UYARI' and '1.234,56' in d['İndirilecek liste ↔ beyan'][1]


def test_yuzde_80_altinda_ekli_sayisi_artar(ay, ayar, tmp_path):
    ted = S.varsayilan_tedarikciler()
    for t in ted:
        if t.vkn != S.ITH:
            t.kdv = 500_000.0
    _, sonuc, _ = _uret(ay, ayar, tmp_path, tedarikciler=ted)
    assert any('EKLİ tutanak sayısı' in k.aciklama for k in sonuc.kontroller)
    assert sonuc.ozet['ekli_toplam'] >= 0.8 * sonuc.ozet['indirilecek_beyan']


def test_410_yoksa_tevkifat_bolumu_isaretlenir(ay, ayar, tmp_path):
    d, sonuc, _ = _uret(ay, ayar, tmp_path, k410=None, i410=None, teminat={'301': 360_000.0})
    uyarilar = [k for k in sonuc.kontroller if k.konu in ('3.7 Tevkifat', 'Geçersiz paragraf')]
    assert any(k.konu == '3.7 Tevkifat' for k in uyarilar)
    assert any('3.7 (410)' in k.aciklama for k in uyarilar)


def test_ocak_sablonundan_subat(tmp_path, ayar):
    eski = S.Senaryo(Donem(2026, 1))
    sab = S.sablon_docx(tmp_path / 'ocak.docx', eski)
    girdi = S.girdi_klasoru(tmp_path / 'G', S.Senaryo(Donem(2026, 2), devreden_onceki=eski.sonraki_devreden))
    out = tmp_path / 't.docx'
    sonuc = taslak_olustur(sab, girdileri_oku(girdi), ayar, out, kati=True)
    metin = _metin(Document(str(out)))
    assert 'sayılı OCAK-2026 dönemi raporumuza eklendiğinden' in metin       # Ocak atfı ŞUBAT'a çevrilmez
    assert 'ŞUBAT/2026' in metin and 'OCAK/2026' not in metin
    assert not [k for k in sonuc.kontroller if k.konu == 'Eski metin' and k.durum == 'UYARI']


def test_araliktan_ocak(tmp_path, ayar):
    eski = S.Senaryo(Donem(2025, 12))
    sab = S.sablon_docx(tmp_path / 'aralik.docx', eski)
    girdi = S.girdi_klasoru(tmp_path / 'G', S.Senaryo(Donem(2026, 1), devreden_onceki=eski.sonraki_devreden))
    out = tmp_path / 't.docx'
    sonuc = taslak_olustur(sab, girdileri_oku(girdi), ayar, out, kati=True)
    metin = _metin(Document(str(out)))
    assert 'OCAK-2026 dönemi raporumuza eklenmiştir.' in metin
    assert 'eklendiğinden bu raporumuza eklenmemiştir' not in metin
    # yıllık güncellenen bölümler (kredi, defter tasdik) program tarafından değiştirilmez, uyarılır
    yil = [k.aciklama for k in sonuc.kontroller if k.konu == 'Yıl ifadesi']
    assert any('"bir önceki yıl" 2024 yazıyor, 2025 olmalı' in a for a in yil)
    assert any('"defter tasdik yılı" 2025 yazıyor, 2026 olmalı' in a for a in yil)
    assert any(k.konu == 'Ocak raporu ekleri' for k in sonuc.kontroller)


def test_ocak_referansi_yoksa_doldurulacak(ay, tmp_path):
    from teminat.ayar import FirmaAyari
    d, sonuc, _ = _uret(ay, FirmaAyari(ymm_no='00000000'), tmp_path)
    p = next(p for p in d.paragraphs if 'raporumuza eklen' in p.text)
    assert '[DOLDURULACAK] tarih ve YMM 00000000/[DOLDURULACAK] sayılı OCAK-2026' in p.text and _sari(p)
    assert any(k.konu == 'Ocak raporu atfı' and k.durum == 'ELLE' for k in sonuc.kontroller)


def test_sablon_ayni_donem_ise_hata(tmp_path, ayar):
    s = S.Senaryo(Donem(2026, 3))
    sab = S.sablon_docx(tmp_path / 'mart.docx', s)
    girdi = S.girdi_klasoru(tmp_path / 'G', s)
    sonuc = taslak_olustur(sab, girdileri_oku(girdi), ayar, tmp_path / 't.docx')
    assert any(k.konu == 'Şablon' and k.durum == 'HATA' for k in sonuc.kontroller)


def test_sablona_yazilmaz(ay, ayar):
    with pytest.raises(SablonHatasi):
        taslak_olustur(ay['sablon'], girdileri_oku(ay['girdi']), ayar, ay['sablon'])


def test_eski_donem_kalan_hucre_uyarilir(ay, ayar, tmp_path):
    d0 = Document(str(ay['sablon']))
    T, _ = tablolari_bul(d0)
    T['tevkifat'].rows[0].cells[0].text = 'ŞUBAT/2026 notu'
    sab2 = tmp_path / 'sablon2.docx'
    d0.save(str(sab2))
    sonuc = taslak_olustur(sab2, girdileri_oku(ay['girdi']), ayar, tmp_path / 't.docx', kati=True)
    assert any(k.konu == 'Eski metin' and k.durum == 'UYARI' and 'ŞUBAT/2026' in k.aciklama for k in sonuc.kontroller)


def test_taslak_uret_dosyalari_ve_uzerine_yazmaz(ay, ayar, tmp_path):
    a = taslak_uret(ayar, ay['sablon'], ay['girdi'])
    assert a.docx.parent == ay['girdi'] / 'CLAUDE TASLAK'
    assert a.docx.name == 'MART-2026 RAPOR TASLAK.docx' and a.xlsx.exists() and a.html.exists()
    b = taslak_uret(ayar, ay['sablon'], ay['girdi'])
    assert b.docx.name == 'MART-2026 RAPOR TASLAK (2).docx' and a.docx.exists()
    from openpyxl import load_workbook
    wb = load_workbook(a.xlsx)
    assert wb.sheetnames == ['Kontrol Listesi', 'Özet', 'Safha (3-4-3)']
    assert wb['Kontrol Listesi'].cell(1, 1).value == 'Durum'
    assert 'Kontrol Listesi' in a.html.read_text(encoding='utf-8')
